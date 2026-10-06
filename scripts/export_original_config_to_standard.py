"""Extract screenshot-sourced activity/coupon config from the original MVP DB.

Reads Zzh032811/ec101-mvp-db (or the local copy of mvp/ec101_mvp.db), maps only
evidenced fields into standard.xlsx sheets, and writes a side-by-side workbook.
Missing source values stay empty. Original columns with no standard counterpart
are listed, not invented.
"""

from __future__ import annotations

import json
import shutil
import sqlite3
import sys
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from mvp.standard_schema import SHEET_COLUMNS, STANDARD_SHEETS
from mvp.standard_workbook import read_standard_workbook, write_standard_workbook

ORIGINAL_DB_CANDIDATES = (
    Path("/tmp/ec101-mvp-db/mvp/ec101_mvp.db"),
    ROOT / "mvp" / "ec101_mvp.db",
)
SAMPLES = ROOT / "docs" / "samples" / "kuaima-verified-standard"
OUT_DIR = SAMPLES / "原库配置对照"
PREVIEW_ROWS = 20

# Converter-produced 活动编号 is the activity name. Original DB has no activity_no.
# activity_import_key is the only evidenced non-name business key (coupon campaigns).
KUAIMA_MANJIAN = "可口可乐满减"
KUAIMA_MANZENG = "满赠优惠"
KUAIMA_TEST1_KEY = "KM-COUPON-AUTO-001"


def _text(value) -> str:
    if value is None:
        return ""
    return str(value)


def _rows(connection: sqlite3.Connection, sql: str, params=()) -> list[dict]:
    connection.row_factory = sqlite3.Row
    return [dict(row) for row in connection.execute(sql, params)]


def load_original(db_path: Path) -> dict[str, list[dict]]:
    connection = sqlite3.connect(db_path)
    dealers = {row["dealer_platform_id"]: row for row in _rows(connection, "SELECT * FROM dealer_platform")}
    activities = _rows(connection, "SELECT * FROM activity ORDER BY activity_id")
    for activity in activities:
        dealer = dealers.get(activity["dealer_platform_id"], {})
        activity["dealer_name"] = dealer.get("dealer_name", "")
        activity["platform_name"] = dealer.get("platform_name", "")
    rules = _rows(connection, "SELECT * FROM activity_rule ORDER BY rule_id")
    benefits = _rows(connection, "SELECT * FROM activity_rule_benefit ORDER BY benefit_id")
    scopes = _rows(connection, "SELECT * FROM activity_scope ORDER BY scope_id")
    issue = _rows(connection, "SELECT * FROM coupon_issue_rule ORDER BY coupon_issue_rule_id")
    use = _rows(connection, "SELECT * FROM coupon_use_rule ORDER BY coupon_use_rule_id")
    connection.close()
    return {
        "activity": activities,
        "activity_rule": rules,
        "activity_rule_benefit": benefits,
        "activity_scope": scopes,
        "coupon_issue_rule": issue,
        "coupon_use_rule": use,
    }


def activity_no(activity: dict) -> str:
    key = activity.get("activity_import_key")
    if key:
        return str(key)
    return str(activity["activity_name"])


def coupon_config_no(activity: dict) -> str:
    return activity_no(activity)


def append_sheet(workbook: Workbook, name: str, header: list[str], rows: list[list[object]]) -> None:
    sheet = workbook.create_sheet(name[:31])
    sheet.append(header)
    for row in rows:
        sheet.append(["" if value is None else value for value in row])
    for cells in sheet.iter_rows():
        for cell in cells:
            cell.number_format = "@"


def map_standard(original: dict[str, list[dict]]) -> tuple[dict[str, list[dict[str, str]]], list[list[str]]]:
    activities = {row["activity_id"]: row for row in original["activity"]}
    rules = {row["rule_id"]: row for row in original["activity_rule"]}
    issue_by_activity = {row["activity_id"]: row for row in original["coupon_issue_rule"]}
    use_by_activity = {row["activity_id"]: row for row in original["coupon_use_rule"]}
    leftover: list[list[str]] = []
    tables = {sheet: [] for sheet in STANDARD_SHEETS if sheet != "导入清单"}

    activity_unmapped_columns = (
        "activity_id", "dealer_platform_id", "tpm_id", "activity_category",
        "product_scope_type", "disabled_product_scope_type", "purchase_limit_type",
        "purchase_limit_value", "customer_scope_type", "disabled_customer_scope_type",
        "use_device", "limit_recharge_gift", "use_scene", "min_sku_count", "rule_version",
    )
    for activity in original["activity"]:
        number = activity_no(activity)
        tables["标准活动"].append({
            "活动编号": number,
            "活动名称": _text(activity["activity_name"]),
            "活动类型": _text(activity["promotion_type"]),
            "促销方式": _text(activity["promo_method"]),
            "开始时间": _text(activity["start_time"]),
            "结束时间": _text(activity["end_time"]),
            "是否允许叠加活动": _text(activity["allow_stack"]),
            "是否允许叠加优惠券": _text(activity["allow_coupon"]),
            "活动状态": _text(activity["activity_status"]),
        })
        for column in activity_unmapped_columns:
            if activity.get(column) not in (None, ""):
                leftover.append(["activity", str(activity["activity_id"]), activity["activity_name"], column, _text(activity[column]), "标准活动无此列"])
        if not activity.get("activity_import_key"):
            leftover.append(["activity", str(activity["activity_id"]), activity["activity_name"], "活动编号", "", "原库无独立活动编号，标准活动编号暂用活动名称"])

        if activity.get("activity_category") == "券类":
            use = use_by_activity.get(activity["activity_id"], {})
            tables["标准优惠券配置"].append({
                "优惠券配置编号": coupon_config_no(activity),
                "优惠券名称": _text(activity["activity_name"]),
                "券类型": _text(use.get("coupon_type")),
                "配置状态": _text(activity["activity_status"]),
            })
            issue = issue_by_activity.get(activity["activity_id"])
            if issue:
                tables["标准优惠券发放规则"].append({
                    "优惠券配置编号": coupon_config_no(activity),
                    "发放方式": _text(issue["issue_mode"]),
                    "发放开始时间": _text(issue["issue_start_at"]),
                    "发放结束时间": _text(issue["issue_end_at"]),
                    "自动发放时间": _text(issue["auto_issue_at"]),
                    "每次发放数量": _text(issue["coupon_qty_per_grant"]),
                    "每客户领取上限": _text(issue["max_claim_per_customer"]),
                    "每日领取上限": _text(issue["daily_claim_limit"]),
                    "规则状态": _text(issue["rule_status"]),
                })
            if use:
                tables["标准优惠券使用规则"].append({
                    "优惠券配置编号": coupon_config_no(activity),
                    "券类型": _text(use["coupon_type"]),
                    "有效期类型": _text(use["validity_mode"]),
                    "使用开始时间": _text(use["use_start_at"]),
                    "使用结束时间": _text(use["use_end_at"]),
                    "领取后有效天数": _text(use["valid_days_after_receive"]),
                    "每张券最多使用次数": _text(use["max_use_per_coupon"]),
                    "规则状态": _text(use["rule_status"]),
                })

    for rule in original["activity_rule"]:
        activity = activities[rule["activity_id"]]
        tables["标准活动规则"].append({
            "活动编号": activity_no(activity),
            "规则编号": _text(rule["tier_no"]),
            "门槛类型": _text(rule["threshold_type"]),
            "门槛值": _text(rule["threshold_value"]),
            "立减金额": _text(rule["reduce_amount"]),
            "是否免邮": _text(rule["free_shipping"]),
        })

    for benefit in original["activity_rule_benefit"]:
        rule = rules[benefit["rule_id"]]
        activity = activities[rule["activity_id"]]
        tables["标准活动权益"].append({
            "活动编号": activity_no(activity),
            "规则编号": _text(rule["tier_no"]),
            "权益编号": _text(benefit["benefit_id"]),
            "权益类型": _text(benefit["benefit_type"]),
            "赠品商品编号": _text(benefit["gift_product_no"]),
            "赠品商品名称": _text(benefit["gift_product_name"]),
            "赠品数量": _text(benefit["gift_qty"]),
            "赠品单位": "",
        })
        for column in ("coupon_id", "coupon_name", "coupon_qty"):
            if benefit.get(column) not in (None, ""):
                leftover.append(["activity_rule_benefit", str(benefit["benefit_id"]), activity["activity_name"], column, _text(benefit[column]), "标准活动权益无送券列，不改写成赠品"])

    for scope in original["activity_scope"]:
        activity = activities[scope["activity_id"]]
        row = {
            "范围编号": _text(scope["scope_id"]),
            "范围类别": _text(scope["scope_category"]),
            "范围维度": _text(scope["scope_dimension"]),
            "范围值": _text(scope["scope_value"]),
        }
        if activity.get("activity_category") == "券类":
            tables["标准优惠券适用范围"].append({"优惠券配置编号": coupon_config_no(activity), **row})
        else:
            tables["标准活动范围"].append({"活动编号": activity_no(activity), **row})

    return tables, leftover


def write_side_by_side(original: dict[str, list[dict]], mapped: dict[str, list[dict[str, str]]], unmapped: list[list[str]], db_path: Path) -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    append_sheet(workbook, "说明", ["项", "值"], [
        ["原库", str(db_path)],
        ["原库来源", "https://github.com/Zzh032811/ec101-mvp-db 的 mvp/ec101_mvp.db"],
        ["规则", "只映射原库里有值、且标准表有对应列的字段；没有的不编"],
        ["活动编号", "原库无活动编号。有 activity_import_key 就用它，否则用活动名称（与快马核销明细一致）"],
        ["券门槛", "原库写在 activity_rule，标准优惠券表没有门槛列，所以门槛落在标准活动规则"],
        ["舟谱", "原库里羿柏/舟谱两条也原样抽出，不并进快马满减/满赠 standard.xlsx"],
    ])
    identity_header = ["原库activity_id", "经销商", "平台", "活动名称", "活动大类", "标准活动编号", "优惠券配置编号", "activity_import_key"]
    identity_rows = []
    for activity in original["activity"]:
        identity_rows.append([
            activity["activity_id"], activity["dealer_name"], activity["platform_name"],
            activity["activity_name"], activity["activity_category"], activity_no(activity),
            coupon_config_no(activity) if activity.get("activity_category") == "券类" else "",
            activity.get("activity_import_key") or "",
        ])
    append_sheet(workbook, "主键对照", identity_header, identity_rows)

    original_sheets = {
        "原库_活动": ("activity", list(original["activity"][0].keys()) if original["activity"] else []),
        "原库_活动规则": ("activity_rule", list(original["activity_rule"][0].keys()) if original["activity_rule"] else []),
        "原库_活动权益": ("activity_rule_benefit", list(original["activity_rule_benefit"][0].keys()) if original["activity_rule_benefit"] else []),
        "原库_活动范围": ("activity_scope", list(original["activity_scope"][0].keys()) if original["activity_scope"] else []),
        "原库_券发放规则": ("coupon_issue_rule", list(original["coupon_issue_rule"][0].keys()) if original["coupon_issue_rule"] else []),
        "原库_券使用规则": ("coupon_use_rule", list(original["coupon_use_rule"][0].keys()) if original["coupon_use_rule"] else []),
    }
    for sheet_name, (key, header) in original_sheets.items():
        rows = original[key]
        append_sheet(workbook, sheet_name, header, [[row.get(column) for column in header] for row in rows])

    config_sheets = (
        "标准活动", "标准活动规则", "标准活动权益", "标准活动范围",
        "标准优惠券配置", "标准优惠券发放规则", "标准优惠券使用规则", "标准优惠券适用范围",
    )
    for sheet in config_sheets:
        columns = list(SHEET_COLUMNS[sheet])
        append_sheet(workbook, sheet, columns, [[row.get(column, "") for column in columns] for row in mapped[sheet]])

    field_map = [
        ["activity.activity_name", "标准活动.活动名称", "原值"],
        ["activity.promotion_type", "标准活动.活动类型", "原值（不是转换器推断的满减/满赠）"],
        ["activity.promo_method", "标准活动.促销方式", "原值"],
        ["activity.start_time / end_time", "标准活动.开始时间 / 结束时间", "原值"],
        ["activity.allow_stack / allow_coupon", "标准活动.是否允许叠加活动 / 优惠券", "原值"],
        ["activity.activity_status", "标准活动.活动状态", "原值"],
        ["activity.activity_import_key 或 activity_name", "标准活动.活动编号", "原库无独立活动编号"],
        ["activity_rule.tier_no", "标准活动规则.规则编号", "原库档位号"],
        ["activity_rule.threshold_type / threshold_value / reduce_amount / free_shipping", "标准活动规则对应列", "原值"],
        ["activity_rule_benefit.benefit_type / gift_*", "标准活动权益对应列", "赠品单位原库没有，留空"],
        ["activity_rule_benefit.coupon_name / coupon_qty", "（无）", "不改写成赠品"],
        ["activity_scope.*", "非券类→标准活动范围；券类→标准优惠券适用范围", "范围编号用 scope_id"],
        ["coupon_issue_rule.*", "标准优惠券发放规则", "发放方式保持 auto_grant / manual_claim"],
        ["coupon_use_rule.*", "标准优惠券使用规则", "有效期类型保持 fixed_period"],
        ["activity（券类）名称/状态 + use_rule.coupon_type", "标准优惠券配置", "原库没有 coupon_config 表"],
    ]
    append_sheet(workbook, "字段映射", ["原库", "标准表", "处理"], field_map)
    append_sheet(workbook, "有值但未落入标准列", ["原库表", "原库行", "活动", "原库字段", "原值", "原因"], unmapped)

    path = OUT_DIR / "原库配置与标准表对照.xlsx"
    workbook.save(path)
    return path


def write_standard_config_only(mapped: dict[str, list[dict[str, str]]]) -> Path:
    tables = {sheet: [] for sheet in STANDARD_SHEETS if sheet != "导入清单"}
    for sheet, rows in mapped.items():
        tables[sheet] = rows
    path = OUT_DIR / "原库配置_standard.xlsx"
    write_standard_workbook(path, tables, {
        "模板版本": "v1",
        "转换工具版本": "v1",
        "经销商名称": "（原库多经销商，见对照.xlsx 主键对照）",
        "平台名称": "（原库含快马与舟谱）",
        "数据开始日期": "",
        "数据结束日期": "",
        "生成时间": "",
    })
    return path


def patch_kuaima_workbook(path: Path, mapped: dict[str, list[dict[str, str]]], activity_names: set[str], coupon_keys: set[str]) -> None:
    workbook = read_standard_workbook(path)
    tables = workbook.tables
    existing_activity = {row["活动编号"]: row for row in tables.get("标准活动", [])}

    def keep_activity(row: dict[str, str]) -> bool:
        return row["活动名称"] in activity_names or row["活动编号"] in activity_names

    def keep_coupon(row: dict[str, str]) -> bool:
        return row.get("优惠券配置编号") in coupon_keys

    filled_activities = []
    for row in mapped["标准活动"]:
        if keep_coupon(row) or row["活动编号"] in coupon_keys:
            filled_activities.append(dict(row))
            continue
        if not keep_activity(row):
            continue
        current = existing_activity.get(row["活动名称"]) or existing_activity.get(row["活动编号"]) or {}
        merged = dict(row)
        if current:
            merged["活动编号"] = current.get("活动编号") or row["活动编号"]
            merged["活动名称"] = current.get("活动名称") or row["活动名称"]
            if current.get("活动类型"):
                merged["活动类型"] = current["活动类型"]
        filled_activities.append(merged)
    if filled_activities:
        tables["标准活动"] = filled_activities

    activity_nos = {row["活动编号"] for row in tables["标准活动"]}
    tables["标准活动规则"] = [row for row in mapped["标准活动规则"] if row["活动编号"] in activity_nos or row["活动编号"] in activity_names]
    for row in tables["标准活动规则"]:
        if row["活动编号"] not in activity_nos:
            for activity in tables["标准活动"]:
                if activity["活动名称"] == row["活动编号"]:
                    row["活动编号"] = activity["活动编号"]
    tables["标准活动权益"] = [row for row in mapped["标准活动权益"] if row["活动编号"] in activity_nos or row["活动编号"] in activity_names]
    for row in tables["标准活动权益"]:
        if row["活动编号"] not in activity_nos:
            for activity in tables["标准活动"]:
                if activity["活动名称"] == row["活动编号"]:
                    row["活动编号"] = activity["活动编号"]
    tables["标准活动范围"] = [row for row in mapped["标准活动范围"] if row["活动编号"] in activity_nos or row["活动编号"] in activity_names]
    for row in tables["标准活动范围"]:
        if row["活动编号"] not in activity_nos:
            for activity in tables["标准活动"]:
                if activity["活动名称"] == row["活动编号"]:
                    row["活动编号"] = activity["活动编号"]

    tables["标准优惠券配置"] = [row for row in mapped["标准优惠券配置"] if keep_coupon(row)]
    tables["标准优惠券发放规则"] = [row for row in mapped["标准优惠券发放规则"] if keep_coupon(row)]
    tables["标准优惠券使用规则"] = [row for row in mapped["标准优惠券使用规则"] if keep_coupon(row)]
    tables["标准优惠券适用范围"] = [row for row in mapped["标准优惠券适用范围"] if keep_coupon(row)]
    write_standard_workbook(path, tables, workbook.manifest)


def write_preview(workbook_path: Path) -> Path:
    workbook = read_standard_workbook(workbook_path)
    preview = Workbook()
    preview.remove(preview.active)
    append_sheet(preview, "说明", ["项", "值"], [
        ["完整文件", str(workbook_path.relative_to(ROOT))],
        ["预览规则", f"大表只保留前 {PREVIEW_ROWS} 行"],
    ])
    append_sheet(preview, "导入清单", list(workbook.manifest.keys()), [list(workbook.manifest.values())])
    for sheet in STANDARD_SHEETS:
        if sheet == "导入清单":
            continue
        columns = list(SHEET_COLUMNS[sheet])
        records = workbook.tables.get(sheet, [])[:PREVIEW_ROWS]
        append_sheet(preview, sheet, columns, [[row.get(column, "") for column in columns] for row in records])
    preview_path = workbook_path.with_name("standard_preview.xlsx")
    preview.save(preview_path)
    return preview_path


def write_markdown(original: dict[str, list[dict]], mapped: dict[str, list[dict[str, str]]], db_path: Path) -> Path:
    lines = [
        "# 原库活动规则 / 优惠券配置 → 标准表",
        "",
        f"数据来源：GitHub [Zzh032811/ec101-mvp-db](https://github.com/Zzh032811/ec101-mvp-db) 的 `mvp/ec101_mvp.db`（截图结构化后写入的活动规则和优惠券配置）。本次读取文件：`{db_path}`。",
        "",
        "原则：**原库有的才写，标准表没有对应列的不编。**",
        "",
        "## 原库有哪些配置",
        "",
        "| activity_id | 经销商 | 平台 | 名称 | 大类 | 标准活动编号 | 优惠券配置编号 |",
        "|---|---|---|---|---|---|---|",
    ]
    for activity in original["activity"]:
        coupon_no = coupon_config_no(activity) if activity.get("activity_category") == "券类" else ""
        lines.append(
            f"| {activity['activity_id']} | {activity['dealer_name']} | {activity['platform_name']} | "
            f"{activity['activity_name']} | {activity['activity_category']} | {activity_no(activity)} | {coupon_no} |"
        )
    lines.extend([
        "",
        "## 映射后的标准表行数",
        "",
        "| 标准工作表 | 行数 |",
        "|---|---:|",
    ])
    for sheet in (
        "标准活动", "标准活动规则", "标准活动权益", "标准活动范围",
        "标准优惠券配置", "标准优惠券发放规则", "标准优惠券使用规则", "标准优惠券适用范围",
    ):
        lines.append(f"| {sheet} | {len(mapped[sheet])} |")
    lines.extend([
        "",
        "## 明确没有、因此留空或未写入的",
        "",
        "- 原库没有活动业务编号（除券的 `activity_import_key`）。满减/满赠的标准「活动编号」= 活动名称。",
        "- 原库没有赠品单位 → `赠品单位` 空。",
        "- 原库权益没有独立「权益编号」→ 暂用 `benefit_id`，对照表里标明这不是业务编号。",
        "- 返券的 `coupon_name` / `coupon_qty` 标准活动权益没有对应列 → 不改写成赠品。",
        "- 标准优惠券表没有门槛/立减列。原库券门槛仍在 `activity_rule`，映射进「标准活动规则」。",
        "- 可口可乐产品288返15元券没有发放规则、使用规则 → 那两张标准表不造行。",
        "",
        "## 文件",
        "",
        "- 两边对照：`原库配置与标准表对照.xlsx`（左原库、右标准、另有未编造说明）",
        "- 仅标准格式：`原库配置_standard.xlsx`",
        "- 已写入快马验收表：`../满减/standard.xlsx`（满减规则 + test1 券配置），`../满赠/standard.xlsx`（满赠规则），`../优惠券/standard.xlsx`（从满减抽出的券域，不含满减核销）。舟谱和「新客户投放」只在对照文件里。",
        "",
    ])
    path = OUT_DIR / "README.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def main() -> None:
    db_path = next(path for path in ORIGINAL_DB_CANDIDATES if path.exists())
    original = load_original(db_path)
    mapped, unmapped = map_standard(original)
    side = write_side_by_side(original, mapped, unmapped, db_path)
    config_only = write_standard_config_only(mapped)
    patch_kuaima_workbook(SAMPLES / "满减" / "standard.xlsx", mapped, {KUAIMA_MANJIAN}, {KUAIMA_TEST1_KEY})
    patch_kuaima_workbook(SAMPLES / "满赠" / "standard.xlsx", mapped, {KUAIMA_MANZENG}, set())
    write_preview(SAMPLES / "满减" / "standard.xlsx")
    write_preview(SAMPLES / "满赠" / "standard.xlsx")
    from scripts.export_coupon_verified_standard import export_coupon_sample
    export_coupon_sample(SAMPLES / "满减" / "standard.xlsx")
    write_markdown(original, mapped, db_path)
    dump = {
        "source_db": str(db_path),
        "activities": [{k: activity[k] for k in ("activity_id", "dealer_name", "platform_name", "activity_name", "activity_category", "promotion_type")} | {"standard_activity_no": activity_no(activity)} for activity in original["activity"]],
        "mapped_counts": {sheet: len(rows) for sheet, rows in mapped.items() if sheet.startswith("标准")},
    }
    (OUT_DIR / "export_summary.json").write_text(json.dumps(dump, ensure_ascii=False, indent=2), encoding="utf-8")
    artifacts = Path("/opt/cursor/artifacts")
    if artifacts.exists():
        shutil.copy2(side, artifacts / "original_config_vs_standard.xlsx")
        shutil.copy2(config_only, artifacts / "original_config_standard.xlsx")
        shutil.copy2(SAMPLES / "满减" / "standard.xlsx", artifacts / "kuaima_manjian_standard.xlsx")
        shutil.copy2(SAMPLES / "满赠" / "standard.xlsx", artifacts / "kuaima_manzeng_standard.xlsx")
    print(json.dumps({"source_db": str(db_path), "side_by_side": str(side), "standard_only": str(config_only), "mapped_counts": dump["mapped_counts"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
