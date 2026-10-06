"""Write Cursor-openable markdown next to the binary standard.xlsx files."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from mvp.standard_schema import STANDARD_SHEETS
from mvp.standard_workbook import read_standard_workbook

SAMPLES = ROOT / "docs" / "samples" / "kuaima-verified-standard"
PREVIEW = 20
FULL_SHEETS = {
    "导入清单", "标准活动", "标准活动规则", "标准活动权益", "标准活动范围",
    "标准优惠券配置", "标准优惠券发放规则", "标准优惠券使用规则", "标准优惠券适用范围",
    "标准优惠券核销明细",
}
LARGE_SHEETS = {"标准客户", "标准商品", "标准订单明细", "标准活动核销明细", "标准履约"}


def md_escape(value: object) -> str:
    text = "" if value is None else str(value)
    return text.replace("|", "\\|").replace("\n", " ")


def table(header: list[str], rows: list[list[object]]) -> str:
    if not header:
        return "_空表_\n"
    lines = [
        "| " + " | ".join(md_escape(column) for column in header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(md_escape(value) for value in row) + " |")
    if not rows:
        lines.append("| " + " | ".join("（空）" for _ in header) + " |")
    return "\n".join(lines) + "\n"


def write_standard_markdown(xlsx_path: Path, title: str) -> Path:
    workbook = read_standard_workbook(xlsx_path)
    parts = [
        f"# {title}",
        "",
        f"Cursor 打不开旁边的 `standard.xlsx`（提示 Binary file is not supported）。本页是同一份数据的可读版。完整 Excel 请从 GitHub 下载，或看 Cursor 产物里的同名 xlsx。",
        "",
        f"源文件：`{xlsx_path.relative_to(ROOT)}`",
        "",
        "## 导入清单",
        "",
        table(list(workbook.manifest.keys()), [list(workbook.manifest.values())]),
        "",
        "## 各表行数",
        "",
        table(["工作表", "行数", "本页展示"], [
            [sheet, (1 if sheet == "导入清单" else len(workbook.tables.get(sheet, []))),
             "全部" if sheet in FULL_SHEETS or sheet == "导入清单" else f"前 {PREVIEW} 行"]
            for sheet in STANDARD_SHEETS
        ]),
    ]
    for sheet in STANDARD_SHEETS:
        if sheet == "导入清单":
            continue
        records = workbook.tables.get(sheet, [])
        columns = list(records[0].keys()) if records else []
        shown = records if sheet in FULL_SHEETS else records[:PREVIEW]
        parts.extend([
            "",
            f"## {sheet}（{len(records)} 行）",
            "",
        ])
        if sheet in LARGE_SHEETS and len(records) > PREVIEW:
            parts.append(f"大表只展示前 {PREVIEW} 行，完整数据在 `standard.xlsx`。")
            parts.append("")
        if not columns:
            parts.append("_无数据_\n")
            continue
        parts.append(table(columns, [[row.get(column, "") for column in columns] for row in shown]))
    out = xlsx_path.with_name("可读.md" if xlsx_path.stem == "standard" else f"{xlsx_path.stem}_可读.md")
    out.write_text("\n".join(parts), encoding="utf-8")
    return out


def write_comparison_markdown(xlsx_path: Path) -> Path:
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    parts = [
        "# 原库配置与标准表对照（可读）",
        "",
        "Cursor 打不开 `.xlsx`。本页把对照工作簿每一张表都写成 Markdown。",
        "",
        f"源文件：`{xlsx_path.relative_to(ROOT)}`",
        "",
    ]
    for name in workbook.sheetnames:
        sheet = workbook[name]
        rows = list(sheet.iter_rows(values_only=True))
        parts.append(f"## {name}")
        parts.append("")
        if not rows:
            parts.append("_空表_\n")
            continue
        header = ["" if value is None else str(value) for value in rows[0]]
        body = [["" if value is None else value for value in row] for row in rows[1:]]
        if len(body) > 80:
            parts.append(f"共 {len(body)} 行，本页展示前 80 行。")
            parts.append("")
            body = body[:80]
        parts.append(table(header, body))
        parts.append("")
    workbook.close()
    out = xlsx_path.with_name("对照_可读.md")
    out.write_text("\n".join(parts), encoding="utf-8")
    return out


def main() -> None:
    written = [
        write_standard_markdown(SAMPLES / "满减" / "standard.xlsx", "快马满减标准表（可读）"),
        write_standard_markdown(SAMPLES / "满赠" / "standard.xlsx", "快马满赠标准表（可读）"),
        write_standard_markdown(SAMPLES / "原库配置对照" / "原库配置_standard.xlsx", "原库配置（标准格式，可读）"),
        write_comparison_markdown(SAMPLES / "原库配置对照" / "原库配置与标准表对照.xlsx"),
    ]
    for path in written:
        print(path.relative_to(ROOT), path.stat().st_size)


if __name__ == "__main__":
    main()
