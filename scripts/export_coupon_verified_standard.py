"""Isolate the verified Kuaima coupon sample from the filled 满减 workbook.

Kuaima did not export a standalone coupon order file. The only test coupon
(test1 / KM-COUPON-AUTO-001) was used on a 满减 order, so converter e2e
attaches 优惠券/活动明细 to the 满减 batch.

This script does not invent rows. It copies the already-patched coupon
config, the one redemption, and that redemption's order/customer/products
out of docs/samples/kuaima-verified-standard/满减/standard.xlsx, dropping
满减 executions so the coupon domain can be imported and recalculated alone.

Do not import this workbook into the shared snapshot DB (ec101_standard.db):
that would double-count the same 200 yuan coupon. Snapshot batch 1 remains
the live 满减+券 batch.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from mvp.standard_schema import SHEET_COLUMNS, STANDARD_SHEETS
from mvp.standard_workbook import StandardWorkbook, read_standard_workbook, write_standard_workbook
from scripts.export_kuaima_verified_standard import SHEET_TO_TABLE
from scripts.export_standard_readable import write_standard_markdown

SAMPLES = ROOT / "docs" / "samples" / "kuaima-verified-standard"
OUT = SAMPLES / "优惠券"
COUPON_ACTIVITY_NO = "KM-COUPON-AUTO-001"
PREVIEW_ROWS = 20
COUPON_SHEETS = (
    "标准优惠券配置",
    "标准优惠券发放规则",
    "标准优惠券使用规则",
    "标准优惠券适用范围",
    "标准优惠券核销明细",
)


def extract_coupon_workbook(manjian: StandardWorkbook) -> StandardWorkbook:
    redemptions = list(manjian.tables.get("标准优惠券核销明细") or [])
    if len(redemptions) != 1:
        raise ValueError(f"满减验收表应有 1 条券核销，实际 {len(redemptions)}")
    order_no = redemptions[0]["订单号"]
    if not order_no:
        raise ValueError("券核销没有订单号，无法抽出独立验收表")
    order_lines = [row for row in manjian.tables["标准订单明细"] if row["单据编号"] == order_no]
    if not order_lines:
        raise ValueError(f"满减验收表找不到券使用订单 {order_no}")
    customer_nos = {row["客户编号"] for row in order_lines if row.get("客户编号")}
    customer_nos.add(redemptions[0]["客户编号"])
    product_nos = {row["商品编号"] for row in order_lines if row.get("商品编号")}
    activities = [row for row in manjian.tables["标准活动"] if row["活动编号"] == COUPON_ACTIVITY_NO]
    if not activities:
        raise ValueError("满减验收表没有 KM-COUPON-AUTO-001，请先跑 export_original_config_to_standard.py")
    tables = {sheet: [] for sheet in STANDARD_SHEETS if sheet != "导入清单"}
    tables["标准客户"] = [row for row in manjian.tables["标准客户"] if row["客户编号"] in customer_nos]
    tables["标准商品"] = [row for row in manjian.tables["标准商品"] if row["商品编号"] in product_nos]
    tables["标准订单明细"] = order_lines
    tables["标准活动"] = activities
    tables["标准活动规则"] = [row for row in manjian.tables["标准活动规则"] if row["活动编号"] == COUPON_ACTIVITY_NO]
    tables["标准活动权益"] = []
    tables["标准活动范围"] = []
    tables["标准活动核销明细"] = []
    for sheet in COUPON_SHEETS:
        tables[sheet] = list(manjian.tables.get(sheet) or [])
    tables["标准履约"] = [row for row in manjian.tables["标准履约"] if row["单据编号"] == order_no]
    manifest = dict(manjian.manifest)
    manifest["数据开始日期"] = "2026-09-09"
    manifest["数据结束日期"] = "2026-09-12"
    return StandardWorkbook(tables=tables, manifest=manifest)


def _append_sheet(workbook: Workbook, name: str, header: list[str], rows: list[list[object]]) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(header)
    for row in rows:
        sheet.append(["" if value is None else value for value in row])
    for cells in sheet.iter_rows():
        for cell in cells:
            cell.number_format = "@"


def write_preview(workbook: StandardWorkbook, path: Path) -> Path:
    preview = Workbook()
    preview.remove(preview.active)
    _append_sheet(preview, "说明", ["项", "值"], [
        ["批次", "优惠券"],
        ["来源", "从已补规则的满减/standard.xlsx 抽出券域，不另造行"],
        ["经销商", workbook.manifest.get("经销商名称", "")],
        ["平台", workbook.manifest.get("平台名称", "")],
        ["数据开始日期", workbook.manifest.get("数据开始日期", "")],
        ["数据结束日期", workbook.manifest.get("数据结束日期", "")],
        ["预览规则", f"大表只保留前 {PREVIEW_ROWS} 行；完整数据在 standard.xlsx"],
        ["勿导入共享库", "ec101_standard.db 的满减批次已含这张券，再导会把 200 算两次"],
    ])
    _append_sheet(preview, "各表行数", ["标准工作表", "行数", "落入数据库表"], [
        [sheet, (1 if sheet == "导入清单" else len(workbook.tables.get(sheet, []))), SHEET_TO_TABLE[sheet]]
        for sheet in STANDARD_SHEETS
    ])
    _append_sheet(preview, "导入清单", list(workbook.manifest.keys()), [list(workbook.manifest.values())])
    for sheet in STANDARD_SHEETS:
        if sheet == "导入清单":
            continue
        columns = list(SHEET_COLUMNS[sheet])
        records = workbook.tables.get(sheet, [])[:PREVIEW_ROWS]
        _append_sheet(preview, sheet, columns, [[row.get(column, "") for column in columns] for row in records])
    path.parent.mkdir(parents=True, exist_ok=True)
    preview.save(path)
    return path


def write_note(workbook: StandardWorkbook, path: Path) -> Path:
    redemption = workbook.tables["标准优惠券核销明细"][0]
    lines = [
        "快马试点优惠券没有独立订单/销售导出，只有一张 test1 活动明细。",
        "本表从已补原库规则的满减 standard.xlsx 抽出：",
        f"- 券核销 {redemption['优惠券编号']} 优惠 {redemption['优惠金额']}",
        f"- 使用订单 {redemption['订单号']}",
        f"- 券配置 {workbook.tables['标准优惠券配置'][0]['优惠券配置编号']}",
        "- 不含满减核销行，费用只应算出券 200",
        "不要把本表再导入 docs/samples/kuaima-verified-standard/ec101_standard.db。",
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")
    (path.parent / "README.md").write_text(
        "\n".join([
            "# 快马优惠券验收表",
            "",
            "**在 Cursor 里请打开 [可读.md](可读.md)，不要打开 `standard.xlsx`。**",
            "",
            "为什么没有像满减/满赠那样的第三份快马源文件批次：试点目录 `快马-兴路强-试点/优惠券/` 只有活动明细，没有订单明细和销售明细。唯一一张 test1 券用在满减订单 `1012420526026091000081` 上。",
            "",
            "本目录不是另造数据。它从已补原库规则的 `../满减/standard.xlsx` 抽出券配置、发放/使用规则、范围、核销，以及那一张使用订单。满减核销行故意不带，用来单独验收券域：费用 200、理论 200、一致 1。",
            "",
            "共享验收库 `../ec101_standard.db` **不要**再导入这份表，满减批次已经含这 200 元。",
            "",
            "重新生成：先有补过规则的满减表，再跑 `python3 scripts/export_coupon_verified_standard.py`（`export_original_config_to_standard.py` 末尾也会调用）。",
            "",
        ]),
        encoding="utf-8",
    )
    return path


def export_coupon_sample(manjian_path: Path | None = None) -> dict[str, object]:
    source = manjian_path or (SAMPLES / "满减" / "standard.xlsx")
    extracted = extract_coupon_workbook(read_standard_workbook(source))
    OUT.mkdir(parents=True, exist_ok=True)
    workbook_path = OUT / "standard.xlsx"
    write_standard_workbook(workbook_path, extracted.tables, extracted.manifest)
    preview_path = write_preview(extracted, OUT / "standard_preview.xlsx")
    note_path = write_note(extracted, OUT / "extraction_report.txt")
    readable = write_standard_markdown(workbook_path, "快马优惠券标准表（可读）")
    counts = {sheet: (1 if sheet == "导入清单" else len(extracted.tables.get(sheet, []))) for sheet in STANDARD_SHEETS}
    summary = {
        "batch": "优惠券",
        "source": str(source.relative_to(ROOT)),
        "workbook": str(workbook_path.relative_to(ROOT)),
        "preview": str(preview_path.relative_to(ROOT)),
        "readable": str(readable.relative_to(ROOT)),
        "report": str(note_path.relative_to(ROOT)),
        "manifest": extracted.manifest,
        "sheet_rows": counts,
        "redemption": extracted.tables["标准优惠券核销明细"][0],
    }
    (OUT / "export_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def main() -> None:
    summary = export_coupon_sample()
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
