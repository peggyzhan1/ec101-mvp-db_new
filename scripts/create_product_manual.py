from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/EC101产品使用说明书.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color="000000", size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Heiti TC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti TC")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color="FFFFFF", size=9)
        set_cell_shading(table.rows[0].cells[i], "1F4E78")
        set_cell_border(table.rows[0].cells[i])
        if widths:
            table.rows[0].cells[i].width = Inches(widths[i])
    for r_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=9)
            set_cell_shading(cells[i], "F3F6F9" if r_index % 2 else "FFFFFF")
            set_cell_border(cells[i])
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbered(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.12)
        p.add_run(f"{index}.  ").bold = True
        p.add_run(item)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Heiti TC"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti TC")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.18
for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)):
    styles[name].font.name = "Heiti TC"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti TC")
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run("EC101 费用运营平台产品使用说明书")
sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(16)
run = sub.add_run("标准数据转换 导入 核算与结果查询")
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(89, 89, 89)
meta = doc.add_paragraph()
meta.add_run("适用版本  ").bold = True
meta.add_run("首期 Windows 转换工具与费用平台导入流程\n")
meta.add_run("适用对象  ").bold = True
meta.add_run("经销商业务人员、费用运营人员、数据管理员")

doc.add_paragraph("本说明书用于指导业务人员将快马、舟谱等平台导出的源文件转换为 EC101 标准数据，并上传至费用运营平台完成校验、导入和自动核算。系统不要求业务人员直接操作数据库；原始文件以 sources.zip 形式归档，平台数据库保存标准业务事实、核算结果和导入审计信息。")

add_heading(doc, "一 产品概览", 1)
doc.add_paragraph("EC101 首期采用“平台转换器 + 标准 Excel + 费用平台导入”的工作方式。不同平台的字段差异由转换器处理，费用平台只接收固定格式的标准 Excel，因此业务人员不需要了解各平台的原始字段结构。")
add_table(doc, ["环节", "使用者", "主要操作", "输出"], [
    ("源数据准备", "业务人员", "从快马或舟谱下载指定期间的数据文件", "平台源文件"),
    ("数据转换", "业务人员", "在 Windows 工具中选择平台、经销商和源文件", "standard.xlsx、sources.zip、转换报告"),
    ("平台导入", "费用运营人员", "上传标准 Excel 和原始文件归档", "导入批次与校验结果"),
    ("自动核算", "系统", "分别汇总活动优惠与优惠券优惠", "核算批次、订单权益和质量问题"),
    ("结果查看", "费用运营人员", "查看活动、订单、异常和结算候选", "费用运营工作台"),
], widths=[1.1, 1.2, 3.0, 1.7])

add_heading(doc, "二 使用前准备", 1)
add_heading(doc, "2.1 文件准备", 2)
add_bullets(doc, [
    "确认经销商名称、平台名称和数据覆盖期间。一次导入对应一个经销商、一个平台和一个完整期间。",
    "从平台重新下载源文件时，尽量保留原始文件名，不要在 Excel 中手工改动编号、单号和券号。",
    "活动配置和优惠券配置如不在平台明细中提供，需要在转换工具中按业务确认结果录入。活动和优惠券分别配置，不能用活动编号替代优惠券编号。",
    "建议将同一批次源文件放在独立文件夹中，便于转换工具打包 sources.zip 并在后续核查时追溯。",
])
add_heading(doc, "2.2 Windows 工具", 2)
doc.add_paragraph("首期桌面工具为 Windows 可执行程序。业务人员无需安装 Python；双击 EC101StandardConverter.exe 即可使用。工具生成的 standard.xlsx 是唯一可上传的数据文件，sources.zip 用于保留原始文件证据。")

add_heading(doc, "三 使用 Windows 转换工具", 1)
add_numbered(doc, [
    "双击打开 EC101StandardConverter.exe。首次使用时确认输出目录可写。",
    "在“平台”中选择“快马”或“舟谱”，在“经销商”中输入标准名称。",
    "分别选择客户、商品、订单、活动、优惠券和履约源文件。没有对应文件时可以留空，但订单导入至少需要客户、商品和订单明细。",
    "选择输出目录。建议每次转换使用新的批次目录，例如“经销商_平台_202601”。",
    "点击“开始转换”，等待状态变为“转换完成”。转换失败时根据弹窗提示补齐文件或修正源数据后重新转换。",
    "检查输出目录中的 standard.xlsx、sources.zip 和 conversion_report.txt。确认报告中没有错误，再进入费用平台上传。",
])
add_heading(doc, "3.1 转换结果说明", 2)
add_table(doc, ["文件", "用途", "是否上传"], [
    ("standard.xlsx", "固定 15 张标准数据工作表，供费用平台解析", "必须"),
    ("sources.zip", "原始平台文件归档，便于审计和重转换", "建议一并上传"),
    ("conversion_report.txt", "记录转换平台、源文件摘要和数量统计", "业务留存"),
], widths=[1.7, 4.0, 1.3])

add_heading(doc, "四 标准 Excel 内容", 1)
doc.add_paragraph("标准 Excel 固定包含以下 15 张工作表。即使本批次没有数据，也保留工作表和表头。编号、单号和优惠券编号按文本处理，避免出现科学计数法或前导零丢失。")
sheet_rows = [
    ("导入清单", "模板版本、转换工具版本、经销商、平台、覆盖期间、生成时间", "每批次"),
    ("标准客户", "客户编号、客户名称、客户类型、客户区域、地址、电话、添加时间、业务员", "主数据"),
    ("标准商品", "商品编号、名称、品牌、目录、条码、规格、基本单位、箱规、每箱基本单位数量", "主数据"),
    ("标准订单明细", "订单、客户、商品、数量、金额、活动编号、优惠券编号、业务员、状态、支付方式等", "交易事实"),
    ("标准活动", "活动基本配置", "活动配置"),
    ("标准活动规则", "门槛、立减、免邮规则", "活动配置"),
    ("标准活动权益", "赠品或权益内容", "活动配置"),
    ("标准活动范围", "活动适用客户、商品或区域范围", "活动配置"),
    ("标准活动核销明细", "活动、客户、订单、商品总金额、优惠金额", "活动执行"),
    ("标准优惠券配置", "优惠券配置编号、名称、类型、状态", "优惠券配置"),
    ("标准优惠券发放规则", "发放方式、发放期间、领取上限、规则状态", "优惠券配置"),
    ("标准优惠券使用规则", "有效期、使用次数、使用时间、规则状态", "优惠券配置"),
    ("标准优惠券适用范围", "优惠券适用客户、商品或区域范围", "优惠券配置"),
    ("标准优惠券核销明细", "优惠券编号、客户、领取时间、使用期限、订单号、优惠金额", "优惠券执行"),
    ("标准履约", "单据编号、下游订单编号、履约状态、出库、完成、结款、退货数量", "履约事实"),
]
add_table(doc, ["工作表", "主要字段", "用途"], sheet_rows, widths=[1.55, 4.6, 1.1])

add_heading(doc, "五 在费用平台导入标准数据", 1)
add_numbered(doc, [
    "启动费用运营平台，进入左侧“数据底座”下的“数据接入”。",
    "点击“标准数据 Excel”选择转换工具生成的 standard.xlsx。",
    "在“原始文件归档”中选择 sources.zip。建议上传归档，以便后续核查来源；没有归档时仍可提交标准 Excel。",
    "点击“开始导入”。页面会显示导入状态；导入成功后系统自动生成核算结果。",
    "成功提示会返回导入批次编号。记录该编号，用于后续问题反馈和历史批次查询。",
])
add_heading(doc, "5.1 导入状态", 2)
add_table(doc, ["状态", "含义", "处理方式"], [
    ("导入中", "系统正在校验并写入当前批次", "等待完成，不要重复提交"),
    ("校验失败", "发现必填字段、重复键或跨表关联错误", "下载或查看错误信息，修正源文件后重新转换"),
    ("已核算", "标准数据已落库，活动和优惠券优惠已分别计算", "进入活动中心、证据与异常或核销与对账查看"),
], widths=[1.3, 3.8, 2.1])

add_heading(doc, "六 费用与促销结果查看", 1)
add_heading(doc, "6.1 活动中心", 2)
doc.add_paragraph("活动中心展示活动配置、活动执行和核算结果。金额型活动以优惠金额展示，赠品型活动以应赠数量和实赠数量展示。活动是否走 TPM 不影响标准数据导入。")
add_heading(doc, "6.2 证据与异常", 2)
doc.add_paragraph("当订单、客户、商品、活动或优惠券之间存在缺失关联时，系统会生成质量问题。警告类问题会阻止进入可提交结算候选，提示类问题需要业务确认。")
add_heading(doc, "6.3 核销与对账", 2)
doc.add_paragraph("可提交金额只来自已通过核验的金额型活动。赠品活动保留数量，不在缺少单位成本时折算成人民币。")
add_heading(doc, "6.4 活动与优惠券的关系", 2)
doc.add_paragraph("同一订单可以同时参与活动并使用优惠券，但两者分别核销：订单连接活动核销明细和优惠券核销明细，系统先分别计算活动优惠与券优惠，再在结果层汇总，不使用活动编号关联优惠券。")

add_heading(doc, "七 常见问题", 1)
faq = [
    ("为什么不能直接上传平台原始文件？", "平台只接收固定标准 Excel。不同平台的字段含义和文件结构不同，先转换可以把差异集中在转换器中，避免费用平台为每个平台维护一套导入逻辑。"),
    ("数据库是否保存原始数据？", "不保存原始平台数据的逐行表。系统保存上传归档的路径、哈希和文件大小；sources.zip 本身用于审计和必要时重转换。"),
    ("舟谱没有下游订单编号怎么办？", "下游订单编号允许为空。履约状态和订单号仍可导入，后续补数时再重新生成标准文件。"),
    ("同一订单既有活动又有优惠券，优惠会重复计算吗？", "不会。活动核销与优惠券核销分别计算，订单总优惠只在结果层相加。"),
    ("导入失败后会不会影响上一批数据？", "不会。导入按批次事务处理，当前批次失败会整体回滚，上一批当前版本保持不变。"),
    ("转换失败后应该怎么处理？", "先查看弹窗和 conversion_report.txt，确认必需源文件、表头和输出目录，再重新选择文件转换。不要手工修改标准 Excel 的表头。"),
]
add_table(doc, ["问题", "说明"], faq, widths=[2.4, 4.8])

add_heading(doc, "八 业务操作检查清单", 1)
add_table(doc, ["检查项", "完成"], [
    ("已确认经销商、平台和数据覆盖期间", "□"),
    ("已下载客户、商品、订单及必要的活动、优惠券、履约文件", "□"),
    ("已使用正确的平台转换器生成 standard.xlsx", "□"),
    ("已检查转换报告，没有阻断性错误", "□"),
    ("已保留并上传 sources.zip", "□"),
    ("平台导入状态为已核算", "□"),
    ("已处理导入错误或质量异常", "□"),
    ("已确认金额活动和赠品活动的结果口径", "□"),
], widths=[6.1, 1.1])

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("EC101 费用运营平台产品使用说明书")
footer.runs[0].font.size = Pt(8)
footer.runs[0].font.color.rgb = RGBColor(128, 128, 128)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
