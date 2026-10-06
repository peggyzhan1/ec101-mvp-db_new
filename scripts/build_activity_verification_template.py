"""Build the activity fee-verification report workbook for review."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[1] / "docs" / "templates" / "活动费用核验报告模版.xlsx"

# 参与订单明细：一行 = 一张参与该活动的订单。计费只认「活动优惠金额」。
ORDER_COLUMNS = [
    ("序号", "now", "报告生成", "本表行号"),
    ("经销商名称", "now", "import_batch", "导入清单"),
    ("平台名称", "now", "import_batch", "导入清单"),
    ("活动编号", "now", "activity", "标准活动"),
    ("活动名称", "now", "activity", "标准活动"),
    ("活动类型", "now", "activity", "满减 / 满赠 / 其他"),
    ("核算日", "now", "calculation_run", "本次核算"),
    ("释放截止时刻", "now", "calculation_run", "核算日前 2 天 0 点"),
    ("单据编号", "now", "order_header", "订单头"),
    ("下单时间", "now", "order_header", "用来判断 T-2"),
    ("客户编号", "now", "order_header", "订单头"),
    ("客户名称", "now", "order_header", "订单头"),
    ("所属业务员", "now", "order_header", "可能多人"),
    ("订单状态", "now", "order_header", "释放条件之一，须为「已完成」"),
    ("订单来源", "now", "order_header", "小程序 / App 等，核验备查"),
    ("订单类型", "now", "order_header", "核验备查"),
    ("支付方式", "now", "order_header", "核验备查，不参与释放"),
    ("活动商品金额", "now", "activity_execution", "该单在本活动核销里的商品总金额"),
    ("订单优惠金额（备查）", "now", "order_line 合计", "订单明细优惠加总，只对照，不计本活动费用"),
    ("活动优惠金额", "now", "activity_execution", "关联了活动才有值；满减计费用这一列"),
    ("优惠券优惠金额", "now", "coupon_redemption", "关联了券才有值，否则空"),
    ("满赠数量", "later", "待核算引擎", "关联满赠才有值；现在空，配置进表后再填"),
    ("是否可释放", "now", "release_candidate", "是 / 否"),
    ("释放结论", "now", "release_candidate", "已完成且达到T-2，或不可释放原因"),
    ("理论权益", "later", "待核算引擎", "配置进标准表并改核算后填写，现在空"),
    ("一致性", "later", "待核算引擎", "理论与活动优惠比对，现在空"),
    ("差异金额", "later", "待核算引擎", "活动优惠 − 理论权益，现在空"),
]

LINE_COLUMNS = [
    ("单据编号", "now", "order_header", "回指参与订单"),
    ("商品编号", "now", "order_line", "订单行"),
    ("商品名称", "now", "order_line", "订单行"),
    ("规格", "now", "order_line", "订单行"),
    ("数量", "now", "order_line", "订单行"),
    ("单位", "now", "order_line", "订单行"),
    ("单价", "now", "order_line", "订单行"),
    ("折前金额", "now", "order_line", "订单行"),
    ("行优惠金额", "now", "order_line", "平台行分摊，只展示，不计入活动费用"),
    ("实付金额", "now", "order_line", "订单行"),
]


HEADER = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=11)
TITLE = Font(name="Microsoft YaHei", bold=True, size=16, color="102A43")
LABEL = Font(name="Microsoft YaHei", bold=True, size=10, color="334E68")
BODY = Font(name="Microsoft YaHei", size=10, color="102A43")
NOTE = Font(name="Microsoft YaHei", size=9, color="627D98")
LATER = Font(name="Microsoft YaHei", size=10, color="9B4D1A")
TEAL = PatternFill("solid", fgColor="0F766E")
LATER_FILL = PatternFill("solid", fgColor="FFF4E6")
ROW_ALT = PatternFill("solid", fgColor="F8FAFC")
YES = PatternFill("solid", fgColor="ECFDF3")
NO = PatternFill("solid", fgColor="FEF3C7")
THIN = Border(
    left=Side(style="thin", color="D9E2EC"),
    right=Side(style="thin", color="D9E2EC"),
    top=Side(style="thin", color="D9E2EC"),
    bottom=Side(style="thin", color="D9E2EC"),
)


def _style_header(ws, row, columns, later_names):
    for index, name in enumerate(columns, start=1):
        cell = ws.cell(row, index, name)
        cell.font = HEADER
        cell.fill = LATER_FILL if name in later_names else TEAL
        if name in later_names:
            cell.font = Font(name="Microsoft YaHei", bold=True, color="9B4D1A", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
    ws.row_dimensions[row].height = 28
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(columns))}{row}"
    ws.freeze_panes = f"A{row + 1}"


def _write_row(ws, row, values, later_indexes=(), yes_no_col=None):
    for index, value in enumerate(values, start=1):
        cell = ws.cell(row, index, value)
        cell.font = LATER if index in later_indexes else BODY
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = THIN
        if row % 2 == 0:
            cell.fill = ROW_ALT
        if index in later_indexes:
            cell.fill = LATER_FILL
        if yes_no_col == index:
            cell.fill = YES if value == "是" else NO
            cell.alignment = Alignment(horizontal="center", vertical="center")


def build() -> Path:
    wb = Workbook()

    cover = wb.active
    cover.title = "活动核验结论"
    cover["A1"] = "EC101 活动费用核验报告（评审模版）"
    cover["A1"].font = TITLE
    cover.merge_cells("A1:D1")
    cover["A2"] = "一份报告只对应一个活动 × 一次核算。下面用兴路强「可口可乐满减」真实数字做样例，橙色列是以后才填的。"
    cover["A2"].font = NOTE
    cover.merge_cells("A2:D2")

    summary = [
        ("经销商名称", "深圳市兴路强商贸有限公司"),
        ("平台名称", "快马"),
        ("活动编号", "可口可乐满减"),
        ("活动名称", "可口可乐满减"),
        ("活动类型", "满减"),
        ("数据期间", "2026-08-31 至 2026-09-17"),
        ("核算日", "2026-09-22"),
        ("释放截止时刻", "2026-09-20 00:00:00"),
        ("导入批次", 1),
        ("核算批次", 1),
        ("参与订单数", 136),
        ("可释放订单数", 133),
        ("待释放订单数", 3),
        ("活动优惠合计（核销）", 2040),
        ("可释放金额", 1995),
        ("待释放金额", 45),
        ("计费口径", "只按本活动核销明细的优惠金额；不按订单行优惠，不用规则重算"),
        ("释放口径", "订单状态=已完成，且下单时间早于释放截止时刻"),
        ("理论权益合计", ""),
        ("一致订单数", ""),
        ("报告生成说明", "评审用样例。下载功能未做；字段以本工作簿为准，等你定稿后再开发。"),
    ]
    cover["A4"] = "项目"
    cover["B4"] = "值"
    cover["C4"] = "阶段"
    for cell in (cover["A4"], cover["B4"], cover["C4"]):
        cell.font = HEADER
        cell.fill = TEAL
        cell.border = THIN
    later_labels = {"理论权益合计", "一致订单数"}
    for index, (label, value) in enumerate(summary, start=5):
        cover.cell(index, 1, label).font = LABEL
        cover.cell(index, 1).border = THIN
        cover.cell(index, 2, value).font = BODY
        cover.cell(index, 2).border = THIN
        stage = "以后（配置进标准表后）" if label in later_labels else "现在就能填"
        cover.cell(index, 3, stage).font = NOTE
        cover.cell(index, 3).border = THIN
        if label in later_labels:
            cover.cell(index, 2).fill = LATER_FILL
            cover.cell(index, 3).fill = LATER_FILL
    cover.column_dimensions["A"].width = 28
    cover.column_dimensions["B"].width = 56
    cover.column_dimensions["C"].width = 28
    cover.column_dimensions["D"].width = 18

    orders = wb.create_sheet("参与订单明细")
    orders["A1"] = "一行一张订单。订单优惠金额只备查；关联活动则填活动优惠金额，关联券则填券优惠，满赠则以后填满赠数量。"
    orders["A1"].font = NOTE
    orders.merge_cells("A1:Y1")
    names = [col[0] for col in ORDER_COLUMNS]
    later_names = {col[0] for col in ORDER_COLUMNS if col[1] == "later"}
    _style_header(orders, 2, names, later_names)
    samples = [
        [1, "深圳市兴路强商贸有限公司", "快马", "可口可乐满减", "可口可乐满减", "满减", "2026-09-22", "2026-09-20 00:00:00",
         "1012420526026083100042", "2026-08-31 11:29:13", "WX-00000000000007507527", "筷乐购",
         "刘亚龙(18819192457),杨景花(13036770153),吴志勇(18126310825)", "已完成", "小程序", "普通订单", "货到付款",
         1126.5, 30, 15, "", "", "是", "已完成且达到T-2", "", "", ""],
        [2, "深圳市兴路强商贸有限公司", "快马", "可口可乐满减", "可口可乐满减", "满减", "2026-09-22", "2026-09-20 00:00:00",
         "1012420526026083100044", "2026-08-31 11:44:54", "WX-00000007505077", "正风超市",
         "吴志勇(18126310825)", "已完成", "小程序", "普通订单", "微信支付",
         331.82, 30, 15, "", "", "是", "已完成且达到T-2", "", "", ""],
        [3, "深圳市兴路强商贸有限公司", "快马", "可口可乐满减", "可口可乐满减", "满减", "2026-09-22", "2026-09-20 00:00:00",
         "1012420526026090500083", "2026-09-05 19:21:59", "WX-00000000000007507752", "美家鑫",
         "刘亚龙(18819192457),吴志勇(18126310825)", "部分发货", "小程序", "普通订单", "微信支付",
         534.43, 15, 15, "", "", "否", "订单状态=部分发货", "", "", ""],
    ]
    later_indexes = {i for i, col in enumerate(ORDER_COLUMNS, start=1) if col[1] == "later"}
    yes_col = names.index("是否可释放") + 1
    for offset, row in enumerate(samples):
        _write_row(orders, 3 + offset, row, later_indexes, yes_col)
    yes_no = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
    orders.add_data_validation(yes_no)
    yes_no.add(f"{get_column_letter(yes_col)}3:{get_column_letter(yes_col)}2000")
    widths = [6, 22, 8, 16, 16, 8, 12, 20, 24, 20, 26, 12, 36, 10, 10, 10, 10, 14, 16, 14, 14, 10, 10, 22, 12, 10, 12]
    for index, width in enumerate(widths, start=1):
        orders.column_dimensions[get_column_letter(index)].width = width
    orders.row_dimensions[3].height = 36
    orders.row_dimensions[4].height = 36
    orders.row_dimensions[5].height = 36

    lines = wb.create_sheet("订单商品行（附录）")
    lines["A1"] = "附录：参与订单的商品行，便于抽查。行优惠金额不是活动费用，不要加总当核销。"
    lines["A1"].font = NOTE
    lines.merge_cells("A1:J1")
    _style_header(lines, 2, [col[0] for col in LINE_COLUMNS], set())
    line_samples = [
        ["1012420526026083100042", "1302198T1533852", "2升雪碧*6", "无", 2, "箱", 29.41, 59.6, 0.79, 58.81],
        ["1012420526026083100042", "199068377", "330ml雪碧摩登罐*24", "无", 2, "箱", 37.38, 75.76, 1.01, 74.75],
        ["1012420526026083100042", "199068371", "330ml可口可乐摩登罐*24", "无", 2, "箱", 37.38, 75.76, 1.01, 74.75],
        ["1012420526026083100042", "199068471", "500ML可口可乐*24", "无", 5, "箱", 46.36, 234.95, 3.13, 231.82],
        ["1012420526026083100042", "199068538", "680ml胶可乐*12", "无", 3, "箱", 32.55, 98.97, 1.32, 97.65],
    ]
    for offset, row in enumerate(line_samples):
        _write_row(lines, 3 + offset, row)
    for index, width in enumerate([24, 18, 28, 8, 8, 8, 10, 12, 12, 12], start=1):
        lines.column_dimensions[get_column_letter(index)].width = width

    fields = wb.create_sheet("字段说明")
    fields["A1"] = "请在本表评审：保留、删掉或改名。改完之前不开发下载。"
    fields["A1"].font = NOTE
    fields.merge_cells("A1:F1")
    field_headers = ["工作表", "字段", "阶段", "来源", "说明", "评审意见（请填）"]
    _style_header(fields, 2, field_headers, set())
    catalog = [("活动核验结论", "见封面各行", "now", "activity_fee_summary + calculation_run", "一个活动一次核算的汇总", "")]
    catalog += [("参与订单明细", name, stage, source, note, "") for name, stage, source, note in ORDER_COLUMNS]
    catalog += [("订单商品行（附录）", name, stage, source, note, "") for name, stage, source, note in LINE_COLUMNS]
    for offset, row in enumerate(catalog):
        values = [row[0], row[1], "现在" if row[2] == "now" else "以后", row[3], row[4], ""]
        _write_row(fields, 3 + offset, values, later_indexes={3} if row[2] == "later" else ())
    for index, width in enumerate([22, 22, 8, 36, 42, 22], start=1):
        fields.column_dimensions[get_column_letter(index)].width = width

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
