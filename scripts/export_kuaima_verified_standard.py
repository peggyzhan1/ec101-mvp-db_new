"""Export the verified Kuaima standard workbooks and a live standard-DB inventory.

Uses the same 兴路强 pilot files as tests/e2e/test_standard_pipeline.py.
The e2e run wrote standard.xlsx into a TemporaryDirectory and deleted it;
this script keeps the conversion output for review.
"""

from __future__ import annotations

import json
import sqlite3
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook

from desktop_converter.app import build_conversion_request, run_conversion
from mvp.standard_schema import SHEET_COLUMNS, STANDARD_SHEETS
from mvp.standard_workbook import read_standard_workbook


PILOT = ROOT / "快马-兴路强-试点"
OUT = ROOT / "docs" / "samples" / "kuaima-verified-standard"
PREVIEW_ROWS = 20

BATCHES = {
    "满减": {
        "客户": [PILOT / "User-202609221722.xlsx"],
        "商品": [PILOT / "Product-202609141043.xlsx"],
        "订单": [PILOT / "满减/快马-兴路强-满减-订单明细-20260831-20260917.xls"],
        "活动": [PILOT / "满减/快马-兴路强-满减-活动明细-20260831-20260917.xls"],
        "优惠券": [PILOT / "优惠券/快马-兴路强-优惠券-test1优惠券-活动明细-20260909-20260912.xls"],
    },
    "满赠": {
        "客户": [PILOT / "User-202609221722.xlsx"],
        "商品": [PILOT / "Product-202609141043.xlsx"],
        "订单": [PILOT / "满赠/快马-兴路强-满赠-订单明细-20260819-20260831.xls"],
        "活动": [PILOT / "满赠/快马-兴路强-满赠-活动明细-20260819-20260831.xls"],
    },
}

SHEET_TO_TABLE = {
    "导入清单": "import_batch / dealer_platform",
    "标准客户": "customer",
    "标准商品": "product",
    "标准订单明细": "order_header + order_line",
    "标准活动": "activity",
    "标准活动规则": "activity_rule",
    "标准活动权益": "activity_benefit",
    "标准活动范围": "activity_scope",
    "标准活动核销明细": "activity_execution",
    "标准优惠券配置": "coupon_config",
    "标准优惠券发放规则": "coupon_issue_rule",
    "标准优惠券使用规则": "coupon_use_rule",
    "标准优惠券适用范围": "coupon_scope",
    "标准优惠券核销明细": "coupon_redemption",
    "标准履约": "fulfillment",
}


def _append_sheet(workbook: Workbook, name: str, header: list[str], rows: list[list[object]]) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(header)
    for row in rows:
        sheet.append(["" if value is None else value for value in row])
    for cells in sheet.iter_rows():
        for cell in cells:
            cell.number_format = "@"


def convert_batch(name: str, files: dict[str, list[Path]]) -> dict[str, object]:
    output_dir = OUT / name
    if output_dir.exists():
        for path in output_dir.iterdir():
            if path.is_file():
                path.unlink()
    request = build_conversion_request("快马", "深圳市兴路强商贸有限公司", files, output_dir)
    result = run_conversion(request)
    workbook = read_standard_workbook(result.workbook_path)
    # sources.zip 只是源文件打包，体积大且仓库里已有试点原件，不保留。
    if result.sources_zip_path.exists():
        result.sources_zip_path.unlink()

    preview = Workbook()
    preview.remove(preview.active)
    _append_sheet(preview, "说明", ["项", "值"], [
        ["批次", name],
        ["转换输出", str(result.workbook_path.relative_to(ROOT))],
        ["经销商", workbook.manifest.get("经销商名称", "")],
        ["平台", workbook.manifest.get("平台名称", "")],
        ["数据开始日期", workbook.manifest.get("数据开始日期", "")],
        ["数据结束日期", workbook.manifest.get("数据结束日期", "")],
        ["预览规则", f"大表只保留前 {PREVIEW_ROWS} 行；完整数据在 standard.xlsx"],
    ])
    _append_sheet(preview, "各表行数", ["标准工作表", "行数", "落入数据库表", "备注"], [
        [sheet, (1 if sheet == "导入清单" else len(workbook.tables.get(sheet, []))), SHEET_TO_TABLE[sheet],
         "完整文件见 standard.xlsx" if (sheet != "导入清单" and len(workbook.tables.get(sheet, [])) > PREVIEW_ROWS) else ""]
        for sheet in STANDARD_SHEETS
    ])
    _append_sheet(preview, "导入清单", list(workbook.manifest.keys()), [list(workbook.manifest.values())])
    for sheet in STANDARD_SHEETS:
        if sheet == "导入清单":
            continue
        columns = list(SHEET_COLUMNS[sheet])
        records = workbook.tables.get(sheet, [])
        shown = records[:PREVIEW_ROWS]
        _append_sheet(preview, sheet, columns, [[row.get(column, "") for column in columns] for row in shown])
    preview_path = output_dir / "standard_preview.xlsx"
    preview.save(preview_path)

    counts = {sheet: (1 if sheet == "导入清单" else len(workbook.tables.get(sheet, []))) for sheet in STANDARD_SHEETS}
    return {
        "batch": name,
        "workbook": str(result.workbook_path.relative_to(ROOT)),
        "preview": str(preview_path.relative_to(ROOT)),
        "report": str(result.report_path.relative_to(ROOT)),
        "manifest": workbook.manifest,
        "sheet_rows": counts,
        "workbook_bytes": result.workbook_path.stat().st_size,
    }


def dump_database(db_path: Path) -> dict[str, object]:
    connection = sqlite3.connect(db_path)
    connection.row_factory = sqlite3.Row
    tables = [row[0] for row in connection.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name"
    )]
    inventory = []
    samples = {}
    columns = {}
    for table in tables:
        info = list(connection.execute(f"PRAGMA table_info({table})"))
        columns[table] = [
            {"name": col["name"], "type": col["type"], "notnull": bool(col["notnull"]), "pk": bool(col["pk"])}
            for col in info
        ]
        count = connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
        sample = connection.execute(f"SELECT * FROM {table} LIMIT 1").fetchone()
        inventory.append({"table": table, "rows": count, "columns": [col["name"] for col in info]})
        samples[table] = dict(sample) if sample else None
    batches = [dict(row) for row in connection.execute("SELECT * FROM import_batch ORDER BY import_batch_id")]
    runs = [dict(row) for row in connection.execute("SELECT * FROM calculation_run ORDER BY calculation_run_id")]
    connection.close()
    return {"db_path": str(db_path.relative_to(ROOT)), "tables": inventory, "columns": columns, "samples": samples, "import_batch": batches, "calculation_run": runs}


def write_db_inventory(snapshot: dict[str, object]) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _append_sheet(workbook, "说明", ["项", "值"], [
        ["库文件", snapshot["db_path"]],
        ["含义", "当前平台读取的标准导入库。业务表按 import_batch 分批；核算表由 calculate_batch 另写。"],
        ["DDL", "mvp/ddl/ec101_standard_sqlite.sql"],
        ["对照文档", "docs/标准数据到标准库对照.md"],
    ])
    _append_sheet(workbook, "导入批次", ["import_batch_id", "dealer_name", "platform_name", "coverage_start", "coverage_end", "status", "is_current"],
                  [[row.get(key) for key in ("import_batch_id", "dealer_name", "platform_name", "coverage_start", "coverage_end", "status", "is_current")]
                   for row in snapshot["import_batch"]])
    _append_sheet(workbook, "核算结果", ["calculation_run_id", "import_batch_id", "calc_date", "activity_benefit", "coupon_benefit", "released_activity_benefit", "released_coupon_benefit", "theoretical_activity_benefit", "theoretical_coupon_benefit", "consistent_orders", "participating_orders", "released_orders", "status"],
                  [[row.get(key) for key in ("calculation_run_id", "import_batch_id", "calc_date", "activity_benefit", "coupon_benefit", "released_activity_benefit", "released_coupon_benefit", "theoretical_activity_benefit", "theoretical_coupon_benefit", "consistent_orders", "participating_orders", "released_orders", "status")]
                   for row in snapshot["calculation_run"]])
    _append_sheet(workbook, "表行数", ["表", "行数", "字段"],
                  [[item["table"], item["rows"], " / ".join(item["columns"])] for item in snapshot["tables"]])
    for table, cols in snapshot["columns"].items():
        sample = snapshot["samples"].get(table)
        body = []
        for col in cols:
            body.append([
                col["name"], col["type"],
                "是" if col["pk"] else "",
                "是" if col["notnull"] else "",
                "（空表）" if sample is None else sample.get(col["name"]),
            ])
        _append_sheet(workbook, table[:31], ["字段", "类型", "主键", "必填", "当前库一条样例"], body)
    path = OUT / "ec101_standard库结构与当前数据.xlsx"
    workbook.save(path)
    return path


def write_readme(summaries: list[dict[str, object]], snapshot: dict[str, object], inventory_xlsx: Path) -> Path:
    lines = [
        "# 快马验收用的转换后标准表",
        "",
        "之前 e2e 验收（`tests/e2e/test_standard_pipeline.py`）把 `standard.xlsx` 写在临时目录，进程结束就删了，所以仓库里看不到转换结果。本目录用**同一套兴路强试点源文件、同一套转换器**重新导出，方便直接打开。",
        "",
        "数据已经进过 `mvp/ec101_standard.db`（两批都是 `calculated`）。标准 Excel 是导入前的中间文件；库是导入后的落库结果。",
        "",
        "重新生成：`python3 scripts/export_kuaima_verified_standard.py`",
        "",
        "## 文件",
        "",
        "| 文件 | 是什么 |",
        "|---|---|",
        "| `满减/standard.xlsx` | 满减批次完整标准工作簿（15 张固定表，可再导入） |",
        "| `满减/standard_preview.xlsx` | 同上，大表只留前 20 行，方便打开看结构 |",
        "| `满减/conversion_report.txt` | 转换器当次汇总 |",
        "| `满赠/standard.xlsx` | 满赠批次完整标准工作簿 |",
        "| `满赠/standard_preview.xlsx` | 满赠预览 |",
        "| `满赠/conversion_report.txt` | 转换器当次汇总 |",
        f"| `{inventory_xlsx.name}` | 当前 `ec101_standard.db` 的表、字段、行数、一条样例 |",
        "",
        "## 满减 / 满赠各表行数",
        "",
        "| 标准工作表 | 满减行数 | 满赠行数 | 落入数据库表 |",
        "|---|---:|---:|---|",
    ]
    manjian = next(item for item in summaries if item["batch"] == "满减")
    manzeng = next(item for item in summaries if item["batch"] == "满赠")
    for sheet in STANDARD_SHEETS:
        lines.append(f"| {sheet} | {manjian['sheet_rows'][sheet]} | {manzeng['sheet_rows'][sheet]} | {SHEET_TO_TABLE[sheet]} |")
    lines.extend([
        "",
        "## 当前库里已经有的批次",
        "",
        "| import_batch_id | 期间 | 状态 | 活动优惠合计 | 券优惠合计 | 可释放活动 | 可释放券 | 参与单 / 可释放单 |",
        "|---|---|---|---:|---:|---:|---:|---|",
    ])
    runs = {row["import_batch_id"]: row for row in snapshot["calculation_run"]}
    for batch in snapshot["import_batch"]:
        run = runs.get(batch["import_batch_id"], {})
        lines.append(
            f"| {batch['import_batch_id']} | {batch.get('coverage_start')} ~ {batch.get('coverage_end')} | {batch['status']} | "
            f"{run.get('activity_benefit', '')} | {run.get('coupon_benefit', '')} | "
            f"{run.get('released_activity_benefit', '')} | {run.get('released_coupon_benefit', '')} | "
            f"{run.get('participating_orders', '')} / {run.get('released_orders', '')} |"
        )
    table_rows = {item["table"]: item["rows"] for item in snapshot["tables"]}
    table_help = [
        ("import_batch", "导入批次", "一次上传一份 standard.xlsx"),
        ("import_file", "源文件归档记录", "sources.zip 元数据"),
        ("import_validation_issue", "导入校验问题", "当前空"),
        ("dealer_platform", "经销商×平台", "兴路强 × 快马 一条"),
        ("customer", "客户", "两个批次各导一次客户主数据，所以行数是 2292×2"),
        ("product", "商品", "同上，两个批次各自一份"),
        ("order_header", "订单头", "满减 1764 + 满赠 1257"),
        ("order_line", "订单行", "满减 36157 + 满赠 26254"),
        ("fulfillment", "履约", "快马目前用订单状态生成，一行对应一单"),
        ("activity", "活动主表", "满减「可口可乐满减」+ 满赠「满赠优惠」"),
        ("activity_rule", "活动规则", "快马自动转换不填"),
        ("activity_benefit", "活动权益", "快马自动转换不填"),
        ("activity_scope", "活动范围", "快马自动转换不填"),
        ("activity_execution", "活动核销", "满减 136 + 满赠 98，费用从这里加"),
        ("coupon_config", "优惠券主表", "快马自动转换不填"),
        ("coupon_issue_rule", "券发放规则", "快马自动转换不填"),
        ("coupon_use_rule", "券使用规则", "快马自动转换不填"),
        ("coupon_scope", "券适用范围", "快马自动转换不填"),
        ("coupon_redemption", "券核销", "满减批次 1 张已使用券"),
        ("calculation_run", "核算批次", "每个导入批次一条 RESULT"),
        ("entitlement_check", "权益核对", "有核销的订单各一条"),
        ("release_candidate", "可释放候选", "T-2 + 已完成 判定"),
        ("activity_fee_summary", "活动费用汇总", "每个活动一条"),
        ("calculation_quality_issue", "核算质量问题", "满减 3 张未完成订单"),
    ]
    lines.extend([
        "",
        "## 当前 `ec101_standard.db` 有哪些表",
        "",
        "库文件不进 git（`.gitignore`），但本机这次验收已经写入。下面是当前行数。",
        "",
        "| 表 | 中文 | 当前行数 | 说明 |",
        "|---|---|---:|---|",
    ])
    for table, title, note in table_help:
        lines.append(f"| `{table}` | {title} | {table_rows.get(table, 0)} | {note} |")
    lines.extend([
        "",
        "库表完整字段见 `mvp/ddl/ec101_standard_sqlite.sql`，Excel 列怎么落到字段见 [`docs/标准数据到标准库对照.md`](../../标准数据到标准库对照.md)。",
        "",
    ])
    path = OUT / "README.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    summaries = [convert_batch(name, files) for name, files in BATCHES.items()]
    db_path = ROOT / "mvp" / "ec101_standard.db"
    snapshot = dump_database(db_path) if db_path.exists() else {"db_path": "mvp/ec101_standard.db", "tables": [], "columns": {}, "samples": {}, "import_batch": [], "calculation_run": []}
    inventory_xlsx = write_db_inventory(snapshot) if snapshot["tables"] else OUT / "ec101_standard库结构与当前数据.xlsx"
    write_readme(summaries, snapshot, inventory_xlsx)
    (OUT / "export_summary.json").write_text(json.dumps({
        "batches": summaries,
        "database": {
            "path": snapshot["db_path"],
            "table_rows": {item["table"]: item["rows"] for item in snapshot["tables"]},
            "import_batch": snapshot["import_batch"],
            "calculation_run": snapshot["calculation_run"],
        },
    }, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(json.dumps({"out": str(OUT), "batches": summaries}, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
