import unittest
from pathlib import Path

DDL_PATH = Path(__file__).resolve().parents[1] / "ddl" / "ec101_standard_sqlite.sql"
IMPORT = Path(__file__).resolve().parents[1] / "import_service.py"
DOC = Path(__file__).resolve().parents[2] / "docs" / "标准数据到标准库对照.md"
LOGIC_DOC = Path(__file__).resolve().parents[2] / "docs" / "快马转换与核算逻辑.md"
PLAN_DOC = Path(__file__).resolve().parents[2] / "docs" / "后续实施计划.md"
REPORT_DOC = Path(__file__).resolve().parents[2] / "docs" / "活动费用核验报告模版说明.md"
REPORT_XLSX = Path(__file__).resolve().parents[2] / "docs" / "templates" / "活动费用核验报告模版.xlsx"
SVG_SCRIPT = Path(__file__).resolve().parents[2] / "scripts" / "build_standard_er_svg.py"


class StandardMappingDocsTests(unittest.TestCase):
    def test_every_imported_sheet_is_documented(self):
        source = IMPORT.read_text(encoding="utf-8")
        doc = DOC.read_text(encoding="utf-8")
        for sheet in (
            "标准客户", "标准商品", "标准订单明细", "标准活动", "标准活动规则", "标准活动权益",
            "标准活动范围", "标准活动核销明细", "标准优惠券配置", "标准优惠券发放规则",
            "标准优惠券使用规则", "标准优惠券适用范围", "标准优惠券核销明细", "标准履约",
        ):
            self.assertIn(f'"{sheet}"', source)
            self.assertIn(sheet, doc)

    def test_er_svg_lists_every_ddl_table(self):
        import importlib.util
        spec = importlib.util.spec_from_file_location("build_standard_er_svg", SVG_SCRIPT)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        svg = module.render().read_text(encoding="utf-8")
        tables = [line.split()[2] for line in DDL_PATH.read_text(encoding="utf-8").splitlines() if line.startswith("CREATE TABLE ")]
        self.assertGreaterEqual(len(tables), 20)
        for table in tables:
            self.assertIn(f">{table}<", svg, table)

    def test_er_png_is_generated_in_repo(self):
        png = Path(__file__).resolve().parents[2] / "docs" / "diagrams" / "ec101-standard-er.png"
        self.assertTrue(png.is_file(), png)
        self.assertGreater(png.stat().st_size, 50_000)

    def test_kuaima_and_result_logic_doc_covers_code_entrypoints(self):
        text = LOGIC_DOC.read_text(encoding="utf-8")
        for needle in (
            "convert_kuaima",
            "calculate_batch",
            "calculation_engine",
            "享受促销政策",
            "已完成且达到T-2",
            "activity_execution",
            "coupon_redemption",
            "1995",
        ):
            self.assertIn(needle, text, needle)

    def test_activity_verification_template_is_reviewable(self):
        text = REPORT_DOC.read_text(encoding="utf-8")
        self.assertIn("活动优惠金额", text)
        self.assertIn("参与订单明细", text)
        self.assertTrue(REPORT_XLSX.is_file(), REPORT_XLSX)
        from openpyxl import load_workbook
        workbook = load_workbook(REPORT_XLSX)
        self.assertEqual(workbook.sheetnames[:3], ["活动核验结论", "参与订单明细", "订单商品行（附录）"])
        headers = [cell.value for cell in workbook["参与订单明细"][2]]
        for name in ("单据编号", "订单优惠金额（备查）", "活动优惠金额", "优惠券优惠金额", "满赠数量", "是否可释放"):
            self.assertIn(name, headers)
        coupon_xlsx = REPORT_XLSX.with_name("活动费用核验报告-快马优惠券样例.xlsx")
        coupon_md = coupon_xlsx.with_suffix(".md")
        self.assertTrue(coupon_xlsx.is_file(), coupon_xlsx)
        self.assertTrue(coupon_md.is_file(), coupon_md)
        coupon_book = load_workbook(coupon_xlsx)
        self.assertEqual(coupon_book.sheetnames[:3], ["活动核验结论", "参与订单明细", "订单商品行（附录）"])
        coupon_headers = [cell.value for cell in coupon_book["参与订单明细"][2]]
        self.assertEqual(coupon_headers, headers)
        order = [cell.value for cell in coupon_book["参与订单明细"][3]]
        self.assertEqual(order[coupon_headers.index("单据编号")], "1012420526026091000081")
        self.assertEqual(order[coupon_headers.index("优惠券优惠金额")], 200)
        self.assertEqual(order[coupon_headers.index("活动优惠金额")], None)
        self.assertEqual(order[coupon_headers.index("理论权益")], 200)
        self.assertEqual(order[coupon_headers.index("一致性")], "一致")

    def test_followup_plan_keeps_fee_from_redemption_and_splits_import(self):
        text = PLAN_DOC.read_text(encoding="utf-8")
        for needle in (
            "Golden Reference",
            "import_service",
            "calculation_engine",
            "SUM 核销",
            "理论权益",
            "1995",
            "舟谱",
        ):
            self.assertIn(needle, text, needle)


if __name__ == "__main__":
    unittest.main()
