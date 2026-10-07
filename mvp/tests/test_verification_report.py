import unittest
from pathlib import Path

from openpyxl import load_workbook

from mvp.verification_report import build_activity_report, build_coupon_report

SNAPSHOT = Path(__file__).resolve().parents[2] / "docs" / "samples" / "kuaima-verified-standard" / "ec101_standard.db"


class VerificationReportTests(unittest.TestCase):
    def test_manjian_report_uses_template_sheets_and_released_1995(self):
        body, filename = build_activity_report(SNAPSHOT, 1, 1)
        self.assertIn("可口可乐满减", filename)
        path = Path(self.id().replace(".", "_") + ".xlsx")
        path.write_bytes(body)
        workbook = load_workbook(path)
        path.unlink()
        self.assertEqual(workbook.sheetnames, ["活动核验结论", "参与订单明细", "订单商品行（附录）"])
        cover = {workbook["活动核验结论"].cell(row, 1).value: workbook["活动核验结论"].cell(row, 2).value for row in range(5, 30)}
        self.assertEqual(cover["活动名称"], "可口可乐满减")
        self.assertEqual(cover["参与订单数"], 136)
        self.assertEqual(cover["可释放订单数"], 133)
        self.assertEqual(cover["活动优惠合计（核销）"], 2040)
        self.assertEqual(cover["可释放金额"], 1995)
        self.assertEqual(cover["理论权益合计"], 2040)
        self.assertEqual(cover["一致订单数"], "136 / 136")
        headers = [cell.value for cell in workbook["参与订单明细"][2]]
        self.assertIn("优惠券优惠金额", headers)
        self.assertEqual(workbook["参与订单明细"].max_row, 138)
        coupon_col = headers.index("优惠券优惠金额") + 1
        coupon_values = [workbook["参与订单明细"].cell(row, coupon_col).value for row in range(3, 139)]
        self.assertIn(200, coupon_values)

    def test_coupon_report_matches_reviewed_sample(self):
        body, filename = build_coupon_report(SNAPSHOT, 1, 1)
        self.assertIn("test1", filename)
        path = Path(self.id().replace(".", "_") + ".xlsx")
        path.write_bytes(body)
        workbook = load_workbook(path)
        path.unlink()
        cover = {workbook["活动核验结论"].cell(row, 1).value: workbook["活动核验结论"].cell(row, 2).value for row in range(5, 35)}
        self.assertEqual(cover["活动编号"], "KM-COUPON-AUTO-001")
        self.assertEqual(cover["优惠券优惠合计"], 200)
        self.assertEqual(cover["可释放金额"], 200)
        self.assertEqual(cover["活动优惠合计（核销）"], None)
        self.assertEqual(cover["理论权益合计"], 200)
        self.assertEqual(cover["一致订单数"], "1 / 1")
        headers = [cell.value for cell in workbook["参与订单明细"][2]]
        row = [workbook["参与订单明细"].cell(3, index).value for index in range(1, len(headers) + 1)]
        values = dict(zip(headers, row))
        self.assertEqual(values["单据编号"], "1012420526026091000081")
        self.assertEqual(values["活动优惠金额"], None)
        self.assertEqual(values["优惠券优惠金额"], 200)
        self.assertEqual(values["理论权益"], 200)
        self.assertEqual(values["一致性"], "一致")
        self.assertEqual(workbook["订单商品行（附录）"].max_row, 4)


if __name__ == "__main__":
    unittest.main()
