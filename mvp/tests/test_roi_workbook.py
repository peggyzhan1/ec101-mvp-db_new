import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from api.server import query_fee_tpm_roi
from mvp.roi_workbook import DETAIL_COLUMNS, SUMMARY_COLUMNS, build_roi_workbook

SNAPSHOT = Path(__file__).resolve().parents[2] / "docs" / "samples" / "kuaima-verified-standard" / "ec101_standard.db"


class RoiWorkbookTests(unittest.TestCase):
    def test_snapshot_workbook_lists_platform_dealer_sales_paid_and_fee(self):
        body, filename = build_roi_workbook(query_fee_tpm_roi(SNAPSHOT, {})["rows"])
        self.assertTrue(filename.endswith(".xlsx"))
        workbook = load_workbook(BytesIO(body))
        self.assertEqual(workbook.sheetnames, ["口径说明", "平台经销商汇总", "活动明细"])
        summary = workbook["平台经销商汇总"]
        self.assertEqual([summary.cell(2, index).value for index in range(1, 7)], list(SUMMARY_COLUMNS))
        dealers = {summary.cell(row, 2).value: [summary.cell(row, index).value for index in range(1, 7)] for row in range(3, 6)}
        self.assertEqual(dealers["深圳市兴路强商贸有限公司"][:3], ["快马", "深圳市兴路强商贸有限公司", 3])
        self.assertEqual(dealers["羿柏"][:3], ["舟谱", "羿柏", 2])
        self.assertEqual(dealers["羿柏"][4], 46548.12)
        self.assertEqual(dealers["羿柏"][5], 555)
        totals = [summary.cell(5, index).value for index in range(1, 7)]
        self.assertEqual(totals[0], "合计")
        self.assertEqual(totals[2], 5)
        self.assertEqual(totals[4], 126387.43)
        self.assertEqual(totals[5], 2750)
        detail = workbook["活动明细"]
        self.assertEqual([detail.cell(2, index).value for index in range(1, 13)], list(DETAIL_COLUMNS))
        by_name = {detail.cell(row, 4).value: [detail.cell(row, index).value for index in range(1, 12)] for row in range(3, 8)}
        self.assertEqual(by_name["可口可乐满减"][:3], ["快马", "深圳市兴路强商贸有限公司", "满减/金额"])
        self.assertEqual(by_name["可口可乐满减"][6:11], [133, 24195, 55021.83, 1995, 27.58])
        self.assertEqual(by_name["可口可乐产品288返15元券"][:3], ["舟谱", "羿柏", "满减/金额"])
        self.assertEqual(by_name["可口可乐产品288返15元券"][6:11], [20, 942, 2388.62, 555, 4.3])
        self.assertEqual(by_name["雪碧系列满100元送抱枕"][7:11], [8157, 44159.5, 0, "满赠只统计金额"])
