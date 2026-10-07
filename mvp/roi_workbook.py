"""Workbook of platform × dealer × activity Coke sales, paid amount, and fee.

Numbers come from the same ROI query used on the fee platform. Missing values stay empty.
"""

from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from mvp.verification_report import content_disposition

HEADER = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=11)
TITLE = Font(name="Microsoft YaHei", bold=True, size=16, color="102A43")
LABEL = Font(name="Microsoft YaHei", bold=True, size=10, color="334E68")
BODY = Font(name="Microsoft YaHei", size=10, color="102A43")
NOTE = Font(name="Microsoft YaHei", size=9, color="627D98")
TEAL = PatternFill("solid", fgColor="0F766E")
ALT = PatternFill("solid", fgColor="F8FAFC")
TOTAL = PatternFill("solid", fgColor="ECFDF3")
THIN = Border(
    left=Side(style="thin", color="D9E2EC"),
    right=Side(style="thin", color="D9E2EC"),
    top=Side(style="thin", color="D9E2EC"),
    bottom=Side(style="thin", color="D9E2EC"),
)
KIND = {"money": "满减/金额", "gift": "满赠", "coupon": "优惠券"}
DETAIL_COLUMNS = (
    "平台", "经销商", "类型", "活动 / 券", "上线开始", "上线结束",
    "已完成参与（含可乐行）", "可乐销量（瓶/罐）", "可乐实付", "核销费用", "ROI", "核算批次",
)
SUMMARY_COLUMNS = ("平台", "经销商", "活动/券条数", "可乐销量（瓶/罐）", "可乐实付", "核销费用")


def _kind(row: dict) -> str:
    return KIND.get(row.get("kind") or "", row.get("kind") or "")


def _roi_cell(row: dict):
    if row.get("periodMissing"):
        return ""
    if row.get("kind") == "gift":
        return "满赠只统计金额"
    if row.get("roi") is None:
        return ""
    return row["roi"]


def _num(row: dict, key: str):
    if row.get("periodMissing"):
        return None
    value = row.get(key)
    return None if value is None else value


def _write_header(sheet: Worksheet, row: int, columns: tuple[str, ...]) -> None:
    for index, name in enumerate(columns, start=1):
        cell = sheet.cell(row, index, name)
        cell.font = HEADER
        cell.fill = TEAL
        cell.border = THIN
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_cells(sheet: Worksheet, row: int, values: list, fill=None) -> None:
    for index, value in enumerate(values, start=1):
        cell = sheet.cell(row, index, "" if value is None else value)
        cell.font = BODY
        cell.border = THIN
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if fill is not None:
            cell.fill = fill


def _cover(sheet: Worksheet, generated_at: str, row_count: int) -> None:
    sheet["A1"] = "EC101 平台 / 经销商 / 活动 销量与费用"
    sheet["A1"].font = TITLE
    sheet.merge_cells("A1:B1")
    sheet["A2"] = "按已确认 ROI 口径从标准库现算。没有的数留空，不用规则金额替换核销费用。"
    sheet["A2"].font = NOTE
    sheet.merge_cells("A2:B2")
    items = [
        ("生成时间", generated_at),
        ("活动/券条数", row_count),
        ("订单", "核销参与单 AND 已完成 AND 下单时间 ∈ 上线期间"),
        ("可乐行", "product.brand = 可口可乐"),
        ("销量", "瓶/罐原值；箱/件 × 箱规（1箱=24瓶 取 24）"),
        ("实付", "可乐行 order_line.paid_amount 合计"),
        ("费用", "同一批已完成参与单的核销优惠合计"),
        ("活动期间", "activity.start_time ~ end_time"),
        ("券期间", "coupon_use_rule.use_start_at ~ use_end_at"),
        ("满赠", "列出销量和实付，不算 ROI"),
        ("计费", "只认核销，不替换 1995 / 200 / 675"),
    ]
    sheet["A4"] = "项目"
    sheet["B4"] = "值"
    for cell in (sheet["A4"], sheet["B4"]):
        cell.font = HEADER
        cell.fill = TEAL
        cell.border = THIN
    for index, (label, value) in enumerate(items, start=5):
        sheet.cell(index, 1, label).font = LABEL
        sheet.cell(index, 1).border = THIN
        sheet.cell(index, 2, value).font = BODY
        sheet.cell(index, 2).border = THIN
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 72


def _summary_rows(rows: list[dict]) -> list[list]:
    grouped: dict[tuple[str, str], dict[str, float]] = defaultdict(
        lambda: {"count": 0, "qty": 0.0, "paid": 0.0, "fee": 0.0}
    )
    for row in rows:
        key = (str(row.get("platform") or ""), str(row.get("dealer") or ""))
        bucket = grouped[key]
        bucket["count"] += 1
        qty = _num(row, "cokeQtyBase")
        paid = _num(row, "cokePaidAmount")
        fee = _num(row, "feeAmount")
        if qty is not None:
            bucket["qty"] += float(qty)
        if paid is not None:
            bucket["paid"] += float(paid)
        if fee is not None:
            bucket["fee"] += float(fee)
    table = []
    for (platform, dealer), bucket in sorted(grouped.items()):
        table.append([
            platform, dealer, int(bucket["count"]),
            round(bucket["qty"], 2), round(bucket["paid"], 2), round(bucket["fee"], 2),
        ])
    total = [
        "合计", "",
        sum(row[2] for row in table),
        round(sum(row[3] for row in table), 2),
        round(sum(row[4] for row in table), 2),
        round(sum(row[5] for row in table), 2),
    ]
    table.append(total)
    return table


def _detail_row(row: dict) -> list:
    if row.get("periodMissing"):
        start, end = "期间缺失", ""
    else:
        start, end = row.get("periodStart") or "", row.get("periodEnd") or ""
    return [
        row.get("platform") or "",
        row.get("dealer") or "",
        _kind(row),
        row.get("name") or "",
        start,
        end,
        _num(row, "completedOrders"),
        _num(row, "cokeQtyBase"),
        _num(row, "cokePaidAmount"),
        _num(row, "feeAmount"),
        _roi_cell(row),
        row.get("calcBatchId") or "",
    ]


def build_roi_workbook(rows: list[dict]) -> tuple[bytes, str]:
    generated_at = datetime.now().isoformat(timespec="seconds")
    workbook = Workbook()
    cover = workbook.active
    cover.title = "口径说明"
    _cover(cover, generated_at, len(rows))

    summary = workbook.create_sheet("平台经销商汇总")
    summary["A1"] = "按平台、经销商汇总可乐销量、实付和核销费用。满赠费用为 0，仍计入条数和销量/实付。"
    summary["A1"].font = NOTE
    summary.merge_cells("A1:F1")
    _write_header(summary, 2, SUMMARY_COLUMNS)
    grouped = _summary_rows(rows)
    for offset, values in enumerate(grouped):
        fill = TOTAL if offset == len(grouped) - 1 else (ALT if offset % 2 else None)
        _write_cells(summary, 3 + offset, values, fill)
    for index, width in enumerate((14, 28, 14, 20, 16, 14), start=1):
        summary.column_dimensions[get_column_letter(index)].width = width

    detail = workbook.create_sheet("活动明细")
    detail["A1"] = "一行一个活动或一张券。销量/实付只计可口可乐；费用是同一批已完成核销参与单的优惠合计。"
    detail["A1"].font = NOTE
    detail.merge_cells("A1:L1")
    _write_header(detail, 2, DETAIL_COLUMNS)
    for offset, row in enumerate(rows):
        _write_cells(detail, 3 + offset, _detail_row(row), ALT if offset % 2 else None)
        detail.row_dimensions[3 + offset].height = 22
    for index, width in enumerate((10, 28, 12, 28, 22, 22, 18, 18, 14, 12, 16, 12), start=1):
        detail.column_dimensions[get_column_letter(index)].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), "EC101-平台经销商活动-销量实付费用.xlsx"


# Re-export so callers do not need verification_report for the Content-Disposition helper.
__all__ = ["DETAIL_COLUMNS", "SUMMARY_COLUMNS", "build_roi_workbook", "content_disposition"]
