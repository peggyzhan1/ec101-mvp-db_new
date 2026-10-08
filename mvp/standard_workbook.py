"""Read and write the fixed standard Excel workbook."""

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import Workbook, load_workbook

from .standard_schema import SHEET_COLUMNS, STANDARD_SHEETS


@dataclass(frozen=True)
class StandardWorkbook:
    tables: dict[str, list[dict[str, str]]]
    manifest: dict[str, str]


def write_standard_workbook(
    path: Path,
    tables: Mapping[str, Sequence[Mapping[str, object]]],
    manifest: Mapping[str, object],
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in STANDARD_SHEETS:
        worksheet = workbook.create_sheet(sheet)
        columns = SHEET_COLUMNS[sheet]
        worksheet.append(list(columns))
        records = manifest if sheet == "导入清单" else tables.get(sheet, ())
        if sheet == "导入清单":
            records = [manifest]
        for record in records:
            worksheet.append(["" if record.get(column) is None else str(record.get(column)) for column in columns])
        for row in worksheet.iter_rows():
            for cell in row:
                cell.number_format = "@"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _excel_text(value: object) -> str | None:
    """Return a cell value Excel can store. Blank cells are omitted.

    openpyxl's write-only mode turns an empty string into ``<c t="inlineStr"/>``
    without the required inline string body. Excel then drops those rows, so a
    sheet that contains blanks opens with only the header.
    """
    if value is None:
        return None
    text = str(value)
    return text if text != "" else None


def _column_name(index: int) -> str:
    name = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        name = chr(65 + remainder) + name
    return name


def _stamp_dimensions(path: Path, row_counts: Sequence[int]) -> None:
    """Write each sheet's used range. Some spreadsheet apps only display that range."""
    import zipfile
    from io import BytesIO

    buffer = BytesIO()
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(buffer, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename.startswith("xl/worksheets/sheet") and info.filename.endswith(".xml"):
                sheet_index = int(info.filename.removeprefix("xl/worksheets/sheet").removesuffix(".xml")) - 1
                columns = len(SHEET_COLUMNS[STANDARD_SHEETS[sheet_index]])
                ref = f"A1:{_column_name(columns)}{row_counts[sheet_index]}"
                payload = payload.replace(b"<sheetData>", f'<dimension ref="{ref}"/><sheetData>'.encode(), 1)
            target.writestr(info, payload)
    path.write_bytes(buffer.getvalue())


def write_standard_workbook_fast(
    path: Path,
    tables: Mapping[str, Sequence[Mapping[str, object]]],
    manifest: Mapping[str, object],
) -> None:
    """Stream a large standard workbook. Values stay text so long identifiers round-trip."""
    workbook = Workbook(write_only=True)
    row_counts: list[int] = []
    for sheet in STANDARD_SHEETS:
        worksheet = workbook.create_sheet(sheet)
        columns = list(SHEET_COLUMNS[sheet])
        worksheet.append(columns)
        records = [manifest] if sheet == "导入清单" else tables.get(sheet, ())
        for record in records:
            worksheet.append([_excel_text(record.get(column)) for column in columns])
        row_counts.append(1 + len(records))
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    _stamp_dimensions(path, row_counts)


def read_standard_workbook(path: Path) -> StandardWorkbook:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if tuple(workbook.sheetnames) != STANDARD_SHEETS:
        raise ValueError("工作表必须按固定标准模板存在且顺序一致")
    tables: dict[str, list[dict[str, str]]] = {}
    manifest: dict[str, str] = {}
    for sheet in STANDARD_SHEETS:
        worksheet = workbook[sheet]
        rows = worksheet.iter_rows(values_only=True)
        header = tuple("" if value is None else str(value) for value in next(rows))
        if header != SHEET_COLUMNS[sheet]:
            raise ValueError(f"{sheet} 表头不符合标准模板")
        records = []
        width = len(header)
        for values in rows:
            padded = list(values) + [None] * max(0, width - len(values))
            record = {column: "" if value is None else str(value) for column, value in zip(header, padded)}
            if any(value.strip() for value in record.values()):
                records.append(record)
        if sheet == "导入清单":
            if len(records) != 1:
                raise ValueError("导入清单必须恰好一行")
            manifest = records[0]
        else:
            tables[sheet] = records
    workbook.close()
    return StandardWorkbook(tables=tables, manifest=manifest)
