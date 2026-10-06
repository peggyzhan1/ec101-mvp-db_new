import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from converters.common import ConversionRequest
from converters.kuaima import convert_kuaima
from mvp.standard_workbook import read_standard_workbook


ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / "快马-兴路强-试点"


class KuaimaConverterTests(unittest.TestCase):
    def test_real_sample_generates_reviewed_standard_workbook(self):
        request = ConversionRequest(
            platform="快马",
            dealer_name="深圳市兴路强商贸有限公司",
            source_paths={
                "customer": [BASE / "User-202609221722.xlsx"],
                "product": [BASE / "Product-202609141043.xlsx"],
                "order_detail": [BASE / "满减/快马-兴路强-满减-订单明细-20260831-20260917.xls"],
                "activity_execution": [BASE / "满减/快马-兴路强-满减-活动明细-20260831-20260917.xls"],
                "coupon_redemption": [BASE / "优惠券/快马-兴路强-优惠券-test1优惠券-活动明细-20260909-20260912.xls"],
            },
            activity_inputs=[{
                "活动编号": "KM-MJ-001", "活动名称": "可口可乐满减", "活动类型": "满减",
                "开始时间": "2026-08-31 10:58:00", "结束时间": "2026-09-17 23:59:00", "活动状态": "已结束",
            }],
            coupon_inputs=[],
            output_dir=Path(tempfile.mkdtemp()),
        )
        result = convert_kuaima(request)
        workbook = read_standard_workbook(result.workbook_path)
        self.assertGreater(result.counts["标准客户"], 0)
        self.assertGreater(result.counts["标准商品"], 0)
        self.assertGreater(result.counts["标准订单明细"], 0)
        self.assertEqual(workbook.tables["标准活动"][0]["活动编号"], "KM-MJ-001")
        self.assertTrue(result.sources_zip_path.exists())

    def test_several_activities_bind_execution_files_and_sum_order_discounts(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            _write_sheet(root / "customers.xlsx", ["客户编号", "客户名称"], [["C1", "客户"]])
            _write_sheet(root / "products.xlsx", ["商品编号", "商品名称"], [["P1", "可乐"], ["G1", "抱枕"]])
            _write_sheet(root / "orders-a.xlsx", ["订单编号", "下单时间", "客户编号", "客户名称", "商品编号", "商品名称", "订货数量", "订货金额", "订单状态"], [
                ["O1", "2026-09-01 10:00:00", "C1", "客户", "P1", "可乐", "1", "300", "已完成"],
                ["O2", "2026-08-20 10:00:00", "C-NEW", "新客户", "G1", "抱枕", "1", "0", "已完成"],
            ])
            _write_sheet(root / "orders-b.xlsx", ["订单编号", "下单时间", "客户编号", "客户名称", "商品编号", "商品名称", "订货数量", "订货金额", "订单状态"], [
                ["O1", "2026-09-01 10:00:00", "C1", "客户", "P1", "可乐", "1", "300", "已完成"],
            ])
            _write_sheet(root / "满减-活动明细.xlsx", ["订单号", "客户", "商品总金额", "促销优惠金额"], [["O1", "客户", "200", "10"], ["O1", "客户", "100", "5"], ["O-MISSING", "客户", "300", "15"]])
            _write_sheet(root / "满赠-活动明细.xlsx", ["订单号", "客户", "商品总金额", "促销优惠金额"], [["O2", "新客户", "100", "0"]])
            _write_sheet(root / "其他-活动明细.xlsx", ["订单号", "客户", "商品总金额", "促销优惠金额"], [["O1", "客户", "300", "99"]])
            result = convert_kuaima(ConversionRequest(
                platform="快马", dealer_name="深圳市兴路强商贸有限公司",
                source_paths={
                    "customer": [root / "customers.xlsx"],
                    "product": [root / "products.xlsx"],
                    "order_detail": [root / "orders-a.xlsx", root / "orders-b.xlsx"],
                    "activity_execution": [root / "满减-活动明细.xlsx", root / "满赠-活动明细.xlsx", root / "其他-活动明细.xlsx"],
                },
                activity_inputs=[
                    {"活动编号": "KM-MJ-001", "活动名称": "可口可乐满减", "活动类型": "满减", "开始时间": "2026-08-31 10:58:00", "结束时间": "2026-09-17 23:59:00", "活动状态": "已结束", "核销文件": "满减-活动明细"},
                    {"活动编号": "KM-MZ-001", "活动名称": "满赠优惠", "活动类型": "满赠", "开始时间": "2026-08-19 15:15:00", "结束时间": "2026-08-31 23:59:00", "活动状态": "已结束", "核销文件": "满赠-活动明细"},
                ],
                coupon_inputs=[], output_dir=root / "out",
            ))
            workbook = read_standard_workbook(result.workbook_path)
            executions = {(row["活动编号"], row["订单号"]): row["优惠金额"] for row in workbook.tables["标准活动核销明细"]}
            self.assertEqual(executions, {("KM-MJ-001", "O1"): "15", ("KM-MZ-001", "O2"): "0"})
            self.assertEqual(sum(row["单据编号"] == "O1" for row in workbook.tables["标准订单明细"]), 1)
            self.assertIn("C-NEW", {row["客户编号"] for row in workbook.tables["标准客户"]})
            self.assertIn("unbound_execution_files=1", result.report_path.read_text(encoding="utf-8"))

    def test_one_activity_still_receives_every_execution_file(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            _write_sheet(root / "orders.xlsx", ["订单编号", "下单时间", "客户编号", "商品编号", "订货数量", "订货金额", "订单状态"], [["O1", "2026-09-01 10:00:00", "C1", "P1", "1", "300", "已完成"]])
            _write_sheet(root / "a.xlsx", ["订单号", "促销优惠金额"], [["O1", "10"]])
            _write_sheet(root / "b.xlsx", ["订单号", "促销优惠金额"], [["O1", "5"]])
            result = convert_kuaima(ConversionRequest(
                platform="快马", dealer_name="深圳市兴路强商贸有限公司",
                source_paths={"order_detail": [root / "orders.xlsx"], "activity_execution": [root / "a.xlsx", root / "b.xlsx"]},
                activity_inputs=[{"活动编号": "KM-MJ-001", "活动名称": "可口可乐满减", "活动类型": "满减", "开始时间": "2026-08-31 10:58:00", "结束时间": "2026-09-17 23:59:00", "活动状态": "已结束"}],
                coupon_inputs=[], output_dir=root / "out",
            ))
            workbook = read_standard_workbook(result.workbook_path)
            self.assertEqual(workbook.tables["标准活动核销明细"][0]["优惠金额"], "15")


def _write_sheet(path: Path, headers: list[str], rows: list[list[str]]) -> None:
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    book.save(path)


if __name__ == "__main__":
    unittest.main()
